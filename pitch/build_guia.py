"""Genera pitch/GUIA_DE_ESTUDIO.docx: guía de estudio del equipo QubitGuard.

Uso, desde la raíz del repositorio:
    python pitch/build_guia.py

Requiere python-docx y matplotlib, además de las dependencias de la aplicación
(qiskit, qiskit-aer, qiskit-ibm-runtime): las salidas de Qiskit que aparecen en
la guía (circuitos, vectores de estado, conteos) se calculan al generar el
documento, no se copian a mano. Las cifras de FakeFez n=3..6 proceden de
docs/IBM_HARDWARE.md (el caso n=6 tarda demasiado para regenerarlo aquí).
"""

from __future__ import annotations

import math
import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import qiskit
import qiskit_aer
import qiskit_ibm_runtime
from qiskit import QuantumCircuit, transpile
from qiskit.quantum_info import Statevector
from qiskit_aer import AerSimulator
from qiskit_ibm_runtime.fake_provider import FakeFez

from app.quantum.experiments import (
    _diffuser,
    _phase_oracle_for_basis_state,
    build_grover_circuit,
    compare,
    grover_curve,
    montecarlo,
    optimal_iterations,
    run_grover,
    theoretical_success,
)

OUT = ROOT / "pitch" / "GUIA_DE_ESTUDIO.docx"
CAPTURAS = ROOT / "pitch" / "capturas"

BODY_FONT = "Calibri"
CODE_FONT = "Courier New"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x66, 0x6D)
H1_COLOR = RGBColor(0x0D, 0x36, 0x6B)
H2_COLOR = RGBColor(0x1C, 0x5C, 0xAB)
H3_COLOR = RGBColor(0x33, 0x3A, 0x40)
TEXT_WIDTH_DXA = 9355  # A4 (11906) menos márgenes de 2,25 cm a cada lado

CALLOUTS = {
    "idea": ("Idea clave", "E8F1FB", "2A78D6"),
    "ejemplo": ("Ejemplo", "E9F6EF", "1BAF7A"),
    "jurado": ("Pregunta del jurado", "FDF0E8", "EB6834"),
    "cuidado": ("Cuidado", "FFF6DD", "C98500"),
}


# --------------------------------------------------------------------------- XML helpers

_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
            "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
            "sectPr", "pPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
              "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
              "tblLook", "tblCaption", "tblDescription"],
}


def _insert_ordered(parent, child) -> None:
    """Inserta child respetando el orden del esquema OOXML (Word rechaza otro orden)."""
    order = _ORDER[parent.tag.split("}")[1]]
    rank = order.index(child.tag.split("}")[1])
    for existing in parent:
        name = existing.tag.split("}")[1]
        if name in order and order.index(name) > rank:
            existing.addprevious(child)
            return
    parent.append(child)

def _set_fonts(rpr_parent, name: str) -> None:
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, "get_or_add_rPr") else rpr_parent
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _shade(element_pr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _insert_ordered(element_pr, shd)


def _para_border(paragraph, color: str = "D0D7DE", sides=("top", "left", "bottom", "right"), size: int = 4, space: int = 4) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        bdr.append(el)
    _insert_ordered(ppr, bdr)


def _field(paragraph, instr: str, placeholder: str = "1") -> None:
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r

    paragraph._p.append(fld("begin"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    r.append(t)
    paragraph._p.append(r)
    paragraph._p.append(fld("separate"))
    run = paragraph.add_run(placeholder)
    paragraph._p.append(fld("end"))
    return run


def _cell_shade(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insert_ordered(tcpr, mar)


def _table_width(table, widths_dxa: list[int]) -> None:
    tbl = table._tbl
    tblpr = tbl.tblPr
    w = tblpr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tblpr.append(w)
    w.set(qn("w:w"), str(sum(widths_dxa)))
    w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    _insert_ordered(tblpr, layout)
    grid = tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_dxa):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def _table_borders(table, color="BFC5CC", inside=True, left_only: tuple[str, int] | None = None) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if left_only:
            if side == "left":
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(left_only[1]))
                el.set(qn("w:color"), left_only[0])
            else:
                el.set(qn("w:val"), "nil")
        elif side.startswith("inside") and not inside:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
        if el.get(qn("w:val")) != "nil":
            el.set(qn("w:space"), "0")
        borders.append(el)
    _insert_ordered(tblpr, borders)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _cant_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


# --------------------------------------------------------------------------- builder

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


class Guide:
    def __init__(self) -> None:
        self.doc = Document()
        self.headings: list[tuple[int, str]] = []
        self.figure_no = 0
        self.table_no = 0
        self._styles()
        self._page()

    # ---- setup
    def _styles(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.size = Pt(11)
        normal.font.color.rgb = INK
        _set_fonts(normal.element, BODY_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.1

        for name, size, color, before, after in (
            ("Heading 1", 22, H1_COLOR, 0, 14),
            ("Heading 2", 15, H2_COLOR, 16, 6),
            ("Heading 3", 12, H3_COLOR, 12, 4),
        ):
            s = st[name]
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = color
            _set_fonts(s.element, BODY_FONT)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
        st["Heading 1"].paragraph_format.page_break_before = True

        lb = st["List Bullet"]
        lb.paragraph_format.space_after = Pt(3)
        _set_fonts(lb.element, BODY_FONT)

    def _page(self) -> None:
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.25)
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.2)
        sec.header_distance = Cm(1.1)
        sec.footer_distance = Cm(1.0)
        sec.different_first_page_header_footer = True

        hp = sec.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run("QubitGuard · Guía de estudio del equipo")
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED
        hp.add_run("\t")
        r2 = hp.add_run("IBM Quantum Hackathon · Reto 5.3 EdTech")
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = MUTED
        tabs = hp.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(16.5), alignment=WD_TAB_ALIGNMENT.RIGHT)
        _para_border(hp, color="D0D7DE", sides=("bottom",), size=4, space=4)

        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for text, instr in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
            r = fp.add_run(text)
            r.font.size = Pt(8.5)
            r.font.color.rgb = MUTED
            run = _field(fp, instr)
            run.font.size = Pt(8.5)
            run.font.color.rgb = MUTED

        # Word actualiza índice y números de página al abrir.
        settings = self.doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        anchor = next((el for el in settings if el.tag.split("}")[1] in later), None)
        if anchor is not None:
            anchor.addprevious(upd)
        else:
            settings.append(upd)
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    # ---- inline text
    def _runs(self, paragraph, text: str, size: float | None = None, color: RGBColor | None = None) -> None:
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                _set_fonts(run._r, CODE_FONT)
                run.font.size = Pt((size or 11) - 1.5)
                run.font.color.rgb = RGBColor(0x0B, 0x3D, 0x6E)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(part)
            if size:
                run.font.size = Pt(size)
            if color is not None and not (part.startswith("`")):
                run.font.color.rgb = color

    # ---- blocks
    def h1(self, text: str) -> None:
        self.headings.append((1, text))
        self.doc.add_heading(text, level=1)

    def h2(self, text: str) -> None:
        self.headings.append((2, text))
        self.doc.add_heading(text, level=2)

    def h3(self, text: str) -> None:
        self.doc.add_heading(text, level=3)

    def p(self, text: str, *, align=None, size=None, color=None, after=None, keep=False):
        para = self.doc.add_paragraph()
        self._runs(para, text, size=size, color=color)
        if align is not None:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if keep:
            para.paragraph_format.keep_with_next = True
        return para

    def bullets(self, items: list[str]) -> None:
        for item in items:
            para = self.doc.add_paragraph(style="List Bullet")
            self._runs(para, item)

    def steps(self, items: list[str]) -> None:
        for i, item in enumerate(items, 1):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(4)
            num = para.add_run(f"{i}. ")
            num.bold = True
            num.font.color.rgb = H2_COLOR
            self._runs(para, item)

    def code(self, text: str, size: float = 8.5, caption: str | None = None) -> None:
        if caption:
            cap = self.p(caption, size=9, color=MUTED, after=2, keep=True)
            cap.runs[0].italic = True
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.15)
            pf.right_indent = Cm(0.15)
            if len(lines) <= 22 and i < len(lines) - 1:
                pf.keep_with_next = True
            _shade(para._p.get_or_add_pPr(), "F4F6F8")
            _para_border(para, color="9FB3C8", sides=("left",), size=18, space=6)
            run = para.add_run(line if line else " ")
            _set_fonts(run._r, CODE_FONT)
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 0.6

    def callout(self, kind: str, body: list[str] | str, title: str | None = None) -> None:
        label, fill, edge = CALLOUTS[kind]
        if isinstance(body, str):
            body = [body]
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_width(table, [TEXT_WIDTH_DXA])
        _table_borders(table, left_only=(edge, 24))
        cell = table.rows[0].cells[0]
        _cell_shade(cell, fill)
        _cell_margins(cell, top=110, bottom=110, left=200, right=200)
        _cant_split(table.rows[0]) if sum(len(b) for b in body) < 900 else None
        head = cell.paragraphs[0]
        head.paragraph_format.space_after = Pt(3)
        head.paragraph_format.keep_with_next = True
        hr = head.add_run(label.upper() + (f" · {title}" if title else ""))
        hr.bold = True
        hr.font.size = Pt(9)
        hr.font.color.rgb = RGBColor.from_string(edge)
        for i, text in enumerate(body):
            para = cell.add_paragraph()
            para.paragraph_format.space_after = Pt(3 if i < len(body) - 1 else 0)
            if text.startswith("- "):
                para.paragraph_format.left_indent = Cm(0.5)
                para.paragraph_format.first_line_indent = Cm(-0.35)
                text = "· " + text[2:]
            self._runs(para, text, size=10.5)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 0.8

    def table(self, headers: list[str], rows: list[list[str]], widths_cm: list[float], size: float = 9.5,
              caption: str | None = None, first_col_bold: bool = False) -> None:
        if caption:
            self.table_no += 1
            cap = self.p(f"Tabla {self.table_no}. {caption}", size=9, color=MUTED, after=3, keep=True)
            cap.runs[0].italic = True
        total = sum(widths_cm)
        widths = [round(w / total * TEXT_WIDTH_DXA) for w in widths_cm]
        widths[-1] += TEXT_WIDTH_DXA - sum(widths)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_width(table, widths)
        _table_borders(table)
        _repeat_header(table.rows[0])
        for c, text in enumerate(headers):
            cell = table.rows[0].cells[c]
            _cell_shade(cell, "DCE6F2")
            _cell_margins(cell, 60, 60, 90, 90)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(text)
            run.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0x0D, 0x2B, 0x52)
        for r, row in enumerate(rows, 1):
            _cant_split(table.rows[r])
            for c, text in enumerate(row):
                cell = table.rows[r].cells[c]
                _cell_margins(cell, 50, 50, 90, 90)
                if r % 2 == 0:
                    _cell_shade(cell, "F7F9FB")
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                self._runs(para, text, size=size)
                if first_col_bold and c == 0:
                    for run in para.runs:
                        run.bold = True
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def figure(self, path: Path, width_cm: float, caption: str) -> None:
        self.figure_no += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.p(f"Figura {self.figure_no}. {caption}", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        for run in cap.runs:
            run.italic = True

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- índice
    def insert_toc(self, anchor) -> None:
        """Campo TOC cuyo resultado inicial es la lista de capítulos: se lee sin actualizar
        campos y Word lo rellena con números de página al abrir (updateFields)."""
        entries = [(lvl, text) for lvl, text in self.headings if text != "Índice"]
        paras = []
        for lvl, text in entries:
            para = OxmlElement("w:p")
            paras.append((para, lvl, text))
        for i, (p_el, lvl, text) in enumerate(paras):
            anchor._p.addprevious(p_el)
        # Construir con la API de python-docx sobre los elementos insertados.
        from docx.text.paragraph import Paragraph

        for i, (p_el, lvl, text) in enumerate(paras):
            para = Paragraph(p_el, anchor._parent)
            pf = para.paragraph_format
            pf.space_after = Pt(4 if lvl == 1 else 1)
            pf.space_before = Pt(6 if lvl == 1 else 0)
            pf.left_indent = Cm(0 if lvl == 1 else 0.7)
            pf.tab_stops.add_tab_stop(Cm(16.4), alignment=WD_TAB_ALIGNMENT.RIGHT, leader=WD_TAB_LEADER.DOTS)
            if i == 0:
                begin = OxmlElement("w:r")
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "begin")
                begin.append(fc)
                p_el.append(begin)
                instr = OxmlElement("w:r")
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                instr.append(it)
                p_el.append(instr)
                sep = OxmlElement("w:r")
                fs = OxmlElement("w:fldChar")
                fs.set(qn("w:fldCharType"), "separate")
                sep.append(fs)
                p_el.append(sep)
            run = para.add_run(text)
            run.font.size = Pt(11 if lvl == 1 else 10)
            run.bold = lvl == 1
            if lvl == 1:
                run.font.color.rgb = H1_COLOR
            if i == len(paras) - 1:
                end = OxmlElement("w:r")
                fe = OxmlElement("w:fldChar")
                fe.set(qn("w:fldCharType"), "end")
                end.append(fe)
                p_el.append(end)


# --------------------------------------------------------------------------- datos reales de Qiskit

def qiskit_facts() -> dict:
    facts: dict = {}
    facts["versions"] = (qiskit.__version__, qiskit_aer.__version__, qiskit_ibm_runtime.__version__, np.__version__)
    circ2, k2 = build_grover_circuit(2, "11")
    facts["draw_n2"] = str(circ2.draw("text", fold=120))
    facts["k_n2"] = k2
    facts["oracle_10"] = str(_phase_oracle_for_basis_state(2, "10").draw("text"))
    facts["diffuser_2"] = str(_diffuser(2).draw("text"))

    qc = QuantumCircuit(2)
    qc.h(range(2))
    s_h = Statevector(qc).data.real
    qc.compose(_phase_oracle_for_basis_state(2, "11"), inplace=True)
    s_o = Statevector(qc).data.real
    qc.compose(_diffuser(2), inplace=True)
    s_d = Statevector(qc).data.real
    facts["sv_steps"] = [np.round(v, 4) + 0.0 for v in (s_h, s_o, s_d)]

    compiled2 = transpile(circ2, basis_gates=["cx", "rz", "sx", "x"], seed_transpiler=7, optimization_level=1)
    facts["n2_depth"] = compiled2.depth()
    facts["n2_2q"] = sum(i.operation.num_qubits == 2 for i in compiled2.data)
    facts["n2_counts"] = AerSimulator(seed_simulator=7).run(compiled2, shots=1024, seed_simulator=7).result().get_counts()

    facts["curve4"] = grover_curve(4)
    facts["ideal4"] = run_grover(4, "1010", 2048, 7, 3, False)
    facts["k0_4"] = run_grover(4, "1010", 1024, 7, 0, False)
    fez = FakeFez()
    facts["fez_name"] = fez.name
    facts["fez_qubits"] = fez.num_qubits
    facts["fez_ops"] = sorted(fez.operation_names)
    facts["mc10"] = montecarlo(10, 5000, 1)
    facts["compare40"] = compare(40)["rows"]
    facts["compare128"] = compare(128)["rows"]
    facts["opt_table"] = [
        (n, 2**n, math.degrees(math.asin(1 / math.sqrt(2**n))), optimal_iterations(n),
         theoretical_success(n, optimal_iterations(n)))
        for n in range(2, 7)
    ]
    return facts


FAKEFEZ = [  # docs/IBM_HARDWARE.md, 24-10-2026 UTC, 1024 shots, seed 7, optimización 3
    (3, "101", 2, 137, 39, 821, 0.801758, 0.945313, 0.125),
    (4, "1010", 3, 448, 111, 678, 0.662109, 0.961319, 0.0625),
    (5, "10101", 4, 1456, 517, 188, 0.183594, 0.999182, 0.03125),
    (6, "101010", 6, 4447, 1706, 17, 0.016602, 0.996586, 0.015625),
]


# --------------------------------------------------------------------------- figuras

BLUE, ORANGE, GRAY, INK_HEX, MUTED_HEX = "#2a78d6", "#eb6834", "#8a9099", "#1f2328", "#5f666d"


def _style_axes(ax) -> None:
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#b9bfc6")
    ax.tick_params(colors=MUTED_HEX, labelsize=9)
    ax.grid(axis="y", color="#e6e9ec", linewidth=0.8)
    ax.set_axisbelow(True)


def fig_geometry(path: Path) -> None:
    n = 4
    theta = math.asin(1 / math.sqrt(2**n))
    fig, ax = plt.subplots(figsize=(6.2, 5.0), dpi=200)
    t = np.linspace(-0.05, math.pi * 0.62, 200)
    ax.plot(np.cos(t), np.sin(t), color="#d3d8de", linewidth=1)
    ax.annotate("", xy=(1.18, 0), xytext=(-0.12, 0), arrowprops=dict(arrowstyle="->", color=MUTED_HEX, lw=1))
    ax.annotate("", xy=(0, 1.18), xytext=(0, -0.12), arrowprops=dict(arrowstyle="->", color=MUTED_HEX, lw=1))
    ax.text(1.2, -0.07, "|s'⟩  (todos los demás)", color=MUTED_HEX, fontsize=9, ha="right", va="top")
    ax.text(0.03, 1.19, "|ω⟩  (estado marcado)", color=MUTED_HEX, fontsize=9, va="bottom")
    shades = ["#9ec5f4", "#6da7ec", "#3987e5", "#184f95"]
    for k in range(4):
        phi = (2 * k + 1) * theta
        x, y = math.cos(phi), math.sin(phi)
        ax.annotate("", xy=(x, y), xytext=(0, 0),
                    arrowprops=dict(arrowstyle="-|>", color=shades[k], lw=2, mutation_scale=12))
        p = math.sin(phi) ** 2
        ax.text(x * 1.07 + (0.02 if k < 3 else -0.02), y * 1.07, f"k={k}: {fmt(math.degrees(phi))}°, P = {fmt(p*100)} %",
                fontsize=8.5, color=INK_HEX, ha="left" if k < 3 else "right", va="center")
    arc = np.linspace(0, theta, 30)
    ax.plot(0.3 * np.cos(arc), 0.3 * np.sin(arc), color=ORANGE, lw=1.4)
    ax.text(0.33, 0.035, "θ = 14,5°", color=ORANGE, fontsize=9)
    arc2 = np.linspace(theta, 3 * theta, 30)
    ax.plot(0.5 * np.cos(arc2), 0.5 * np.sin(arc2), color=ORANGE, lw=1.4, linestyle="--")
    ax.text(0.5 * math.cos(2 * theta) + 0.02, 0.5 * math.sin(2 * theta), "2θ por iteración", color=ORANGE, fontsize=8.5)
    ax.set_xlim(-0.45, 1.3)
    ax.set_ylim(-0.15, 1.3)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("Grover como rotación (n = 4, N = 16)", fontsize=11, color=INK_HEX, loc="left")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_curve(path: Path, curve: dict) -> None:
    ks = curve["iterations"]
    probs = curve["probability"]
    N = 2 ** curve["n_qubits"]
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
    _style_axes(ax)
    ax.plot(ks, [p * 100 for p in probs], color=BLUE, lw=2, marker="o", ms=6, zorder=3)
    ax.plot(ks, [min(1, k / N) * 100 for k in ks], color=GRAY, lw=2, linestyle="--", marker="s", ms=5, zorder=2)
    for k, p in zip(ks, probs):
        ax.text(k, p * 100 + 4, f"{fmt(p*100)} %", ha="center", fontsize=8, color=INK_HEX)
    ax.text(0.12, 76, "Grover (Statevector)", color=BLUE, fontsize=9, ha="left")
    ax.text(6.1, 6 / N * 100 + 4, "Azar con k consultas (k/N)", color=MUTED_HEX, fontsize=9, ha="right")
    ax.axvline(curve["optimal_iterations"], color="#c3c8ce", lw=1, linestyle=":")
    ax.text(curve["optimal_iterations"] + 0.08, 104, "k óptimo = 3", fontsize=8.5, color=MUTED_HEX)
    ax.set_xlabel("Iteraciones de Grover (k)", fontsize=9, color=MUTED_HEX)
    ax.set_ylabel("Probabilidad de medir el marcado (%)", fontsize=9, color=MUTED_HEX)
    ax.set_ylim(0, 112)
    ax.set_xticks(ks)
    ax.set_title("n = 4 qubits: la probabilidad sube, llega al máximo y vuelve a bajar", fontsize=10.5,
                 color=INK_HEX, loc="left")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_noise(path: Path) -> None:
    ns = [row[0] for row in FAKEFEZ]
    ideal = [row[7] * 100 for row in FAKEFEZ]
    measured = [row[6] * 100 for row in FAKEFEZ]
    rand = [row[8] * 100 for row in FAKEFEZ]
    depth = [row[3] for row in FAKEFEZ]
    x = np.arange(len(ns))
    w = 0.36
    fig, ax = plt.subplots(figsize=(6.4, 3.7), dpi=200)
    _style_axes(ax)
    ax.bar(x - w / 2 - 0.01, ideal, w, color=BLUE, label="Ideal (teoría)", zorder=3)
    ax.bar(x + w / 2 + 0.01, measured, w, color=ORANGE, label="FakeFez (simulación con ruido)", zorder=3)
    for i in range(len(ns)):
        ax.plot([x[i] - 0.42, x[i] + 0.42], [rand[i]] * 2, color=INK_HEX, lw=1.4, linestyle="--", zorder=4)
        ax.text(x[i] - w / 2, ideal[i] + 2, f"{fmt(ideal[i])}", ha="center", fontsize=8, color=INK_HEX)
        ax.text(x[i] + w / 2, measured[i] + 2, f"{fmt(measured[i])}", ha="center", fontsize=8, color=INK_HEX)
    ax.plot([], [], color=INK_HEX, lw=1.4, linestyle="--", label="Azar 1/N")
    ax.set_xticks(x, [f"n = {n}\nprofundidad {d}" for n, d in zip(ns, depth)])
    ax.set_ylabel("Éxito (%)", fontsize=9, color=MUTED_HEX)
    ax.set_ylim(0, 124)
    ax.legend(frameon=False, fontsize=8.5, loc="upper right", ncol=3, bbox_to_anchor=(1.0, 1.02))
    ax.set_title("Más qubits, circuito más profundo, más ruido acumulado", fontsize=10.5, color=INK_HEX,
                 loc="left", pad=16)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


# --------------------------------------------------------------------------- contenido

def fmt(x: float, d: int = 1) -> str:
    return f"{x:.{d}f}".replace(".", ",")


def cover(g: Guide, f: dict) -> None:
    doc = g.doc
    for _ in range(5):
        doc.add_paragraph()
    g.p("QUBITGUARD QUANTUM LAB", size=12, color=H2_COLOR, after=4).runs[0].bold = True
    t = g.p("Guía de estudio del equipo", size=32, color=H1_COLOR, after=6)
    t.runs[0].bold = True
    g.p("Qué hace el software, cómo funciona por dentro, qué usamos de Qiskit y cómo se "
        "ejecutaría en una computadora cuántica real de IBM", size=14, color=MUTED, after=24)
    bar = doc.add_paragraph()
    _para_border(bar, color="2A78D6", sides=("top",), size=18, space=1)
    g.p("IBM Quantum Hackathon · Qiskit Fall Fest 2026 (Lima)", size=11, after=2)
    g.p("Reto 5.3 EdTech, IEEE CS PUCP", size=11, after=2)
    g.p("Rama `feat/motion-ibm-casaos` · estado del repositorio al 24 de octubre de 2026", size=11, after=2)
    v = f["versions"]
    g.p(f"Salidas de Qiskit generadas con qiskit {v[0]}, qiskit-aer {v[1]}, qiskit-ibm-runtime {v[2]}, NumPy {v[3]}",
        size=10, color=MUTED, after=30)
    g.callout("cuidado", [
        "**Estado honesto que todo el equipo debe repetir igual:** todos los circuitos de QubitGuard se "
        "ejecutaron en simuladores locales (Qiskit Aer, con el modelo de ruido FakeFez como opción). "
        "**No se envió ningún job a hardware cuántico real de IBM.** El script para hacerlo existe y está "
        "probado con dobles locales, pero no hay ejecución real. Tampoco se hizo todavía un piloto con estudiantes.",
    ], title="leer antes que nada")


def how_to_use(g: Guide) -> None:
    g.h1("Cómo usar esta guía")
    g.p("Esta guía está escrita para que cualquier integrante del equipo, venga de la carrera que venga y esté en el "
        "ciclo que esté, pueda explicar QubitGuard ante el jurado y defenderlo con datos. No hace falta saber física "
        "cuántica ni programación avanzada: cada concepto se construye desde cero y siempre se conecta con el código real "
        "del repositorio.")
    g.p("Todo lo que aparece aquí sale de los archivos del proyecto (README, carpeta `docs/`, código en `app/`, "
        "`frontend/src/` y `scripts/`). Las salidas de Qiskit (dibujos de circuitos, vectores de estado, conteos) "
        "se calcularon de verdad al generar este documento. Cuando una cifra es una simulación, se dice.")
    g.h3("Los recuadros")
    g.table(["Recuadro", "Para qué sirve"], [
        ["**Idea clave**", "Lo mínimo que debes poder decir con tus palabras. Si solo lees estos recuadros, ya puedes defender lo esencial."],
        ["**Ejemplo**", "Un caso concreto con números, casi siempre calculado con el propio código del proyecto."],
        ["**Pregunta del jurado**", "Una pregunta que probablemente nos hagan, con una respuesta modelo corta."],
        ["**Cuidado**", "Afirmaciones que NO debemos hacer porque serían falsas o exageradas."],
    ], [4, 12])
    g.h3("Ruta de lectura sugerida")
    g.table(["Si tienes...", "Lee"], [
        ["20 minutos", "Capítulo 1, los recuadros de Idea clave de los capítulos 2 a 6 y el capítulo 8 (preguntas del jurado)."],
        ["1 hora", "Capítulos 1, 2, 4 y 8, y la sección 6.6 (qué esperar en hardware real)."],
        ["2 horas", "Todo, en orden. Al final, intenta responder las preguntas del capítulo 8 sin mirar."],
    ], [4, 12])
    g.p("Convención: las fórmulas se escriben en texto plano. `N = 2^n` significa «N es 2 elevado a n»; "
        "`√N` es la raíz cuadrada de N; `sin²(x)` es el seno de x al cuadrado; θ (theta) y φ (phi) son ángulos.")
    g.h2("Índice")
    g.toc_anchor = g.doc.add_paragraph()


def chapter1(g: Guide, f: dict) -> None:
    g.h1("1. Qué es QubitGuard y qué problema resuelve")
    g.h2("1.1 En una frase")
    g.p("**QubitGuard es un laboratorio web educativo de ciberseguridad cuántica.** El estudiante **predice** una "
        "probabilidad o un número de iteraciones, **ejecuta** un circuito del algoritmo de Grover de 2 a 6 qubits "
        "construido con Qiskit, **compara** su predicción con la teoría y con lo medido en un simulador (ideal o con "
        "el ruido de un chip de IBM simulado, FakeFez) y **practica** cada día con un reto breve.")
    g.p("La historia que conecta todo es una pregunta sencilla: **¿cuántos intentos hacen falta para encontrar un "
        "secreto?** Adivinando al azar sin repetir, se esperan `(N+1)/2` intentos. El algoritmo de Grover "
        "amplifica la amplitud del estado correcto y necesita del orden de `π/4 · √N` iteraciones.")
    g.callout("idea", [
        "QubitGuard enseña **cómo cambia una distribución de probabilidad cuando se manipulan amplitudes**. "
        "La búsqueda de un secreto es solo la excusa para que eso se vea, se prediga y se mida.",
    ])
    if (CAPTURAS / "02_inicio.png").exists():
        g.figure(CAPTURAS / "02_inicio.png", 13, "Pantalla de inicio de QubitGuard (captura real de la aplicación).")

    g.h2("1.2 El problema que atacamos")
    g.bullets([
        "**La probabilidad cuántica no se deja intuir.** Las amplitudes pueden ser negativas y cancelarse entre sí; "
        "nada en la estadística escolar lo anticipa. Una fórmula como `sin²((2k+1)θ)` no se aprende leyéndola: hay "
        "que verla moverse.",
        "**La amenaza cuántica a la criptografía se cuenta mal.** Se suele reducir a «romperá todas las contraseñas». "
        "La realidad es más matizada: Grover da una ventaja solo cuadrática (duplicar los bits de la clave restaura el "
        "margen) y el riesgo concreto es para RSA y curvas elípticas por el algoritmo de Shor.",
        "**Falta práctica activa y breve.** Muy pocas herramientas piden al estudiante comprometerse con una "
        "predicción antes de ver el resultado.",
    ])

    g.h2("1.3 El reto 5.3 EdTech y sus pares oficiales")
    g.p("La página oficial de retos propone «pares» entre un concepto clásico y su equivalente cuántico. Un proyecto "
        "EdTech debe enseñar alguno de ellos.")
    g.table(["Concepto clásico", "Equivalente cuántico sugerido", "¿Lo cubre QubitGuard?"], [
        ["Distribuciones de probabilidad", "Superposición y amplitudes", "**Sí: es el eje principal.**"],
        ["Monte Carlo", "Quantum Amplitude Estimation (QAE)", "Solo la parte clásica, como referencia de coste. **No implementamos QAE.**"],
        ["SVM", "QSVM y kernels cuánticos", "No."],
        ["Redes neuronales", "QNN y VQC", "No."],
        ["Descenso de gradiente", "Quantum Gradient Descent", "No."],
    ], [4.2, 5, 6.8], caption="Pares sugeridos por la página oficial del reto (docs/RETO_Y_ENTREGA.md).")
    g.p("**Par principal: distribuciones de probabilidad → superposición y amplitudes.** Con cero iteraciones, el "
        "circuito produce una distribución uniforme (todos los resultados igual de probables). Al aplicar el oráculo y "
        "el difusor, Grover concentra la probabilidad en el estado marcado. Con ruido, se discute la diferencia entre "
        "lo ideal y lo observado.")
    g.table(["Lo que dice la estadística", "Lo que añade la cuántica (en QubitGuard)"], [
        ["Distribución uniforme sobre `N = 2^n` resultados, `p = 1/N`.", "Superposición uniforme tras las puertas Hadamard: cada estado tiene amplitud `a = 1/√N` y `p = a²`."],
        ["Las probabilidades son ≥ 0 y suman 1.", "Las amplitudes pueden ser negativas; la probabilidad es `a²` y la suma de los cuadrados es 1."],
        ["Una distribución uniforme no favorece a nadie.", "El oráculo cambia el signo del marcado y el difusor refleja todo respecto a la media: la amplitud del marcado sube."],
        ["Más muestras, más información.", "Con k iteraciones, `P(k) = sin²((2k+1)θ)`; pasarse del óptimo **reduce** la probabilidad."],
    ], [7, 9], caption="De la estadística a la amplitud.")
    g.p("**Segundo par, tocado solo de lado: Monte Carlo.** En la app, «Monte Carlo» es la **búsqueda aleatoria "
        "clásica sin repetición** de un secreto, simulada con NumPy. Sirve como referencia para medir la ventaja de "
        "Grover. La estimación cuántica de amplitud (QAE) es otro algoritmo: estima una probabilidad con menos muestras "
        "que un Monte Carlo clásico. Sería un alcance nuevo, con diseño y pruebas propios, y no lo fingimos.")
    g.callout("cuidado", [
        "Nunca digas «implementamos QAE» ni «hacemos Monte Carlo cuántico». Lo correcto es: «el par que cubrimos es "
        "distribuciones → amplitudes; el Monte Carlo de la app es clásico y sirve de comparación con Grover».",
    ])

    g.h2("1.4 Para quién, con qué frecuencia y por qué volver")
    g.table(["Pregunta de las bases", "Respuesta del proyecto", "Estado"], [
        ["¿Desde qué edad?", "**Desde 14 años.** Solo exige fracciones, porcentajes y la idea de probabilidad; no exige álgebra lineal.", "Público propuesto; no validado con menores."],
        ["¿Uso diario o semanal?", "**Diario, unos 3 minutos:** un reto del día, una predicción y un experimento. Para un docente, también sirve como actividad semanal de aula.", "Diseño; la retención no se ha medido."],
        ["¿Metodología?", "**Predice, Observa, Explica (POE):** el estudiante apuesta antes de mirar, observa teoría y medición, y lee una explicación que corrige su error.", "Implementada."],
        ["¿Razón para volver?", "Reto diario determinista por fecha, racha, XP, 7 niveles, 5 pasos de ruta y 6 logros.", "Implementada; su efecto en el hábito no está probado."],
    ], [3.6, 8.4, 4], caption="Público y uso (README, sección 1).")

    g.h2("1.5 Lo que QubitGuard NO es")
    g.p("No es un gestor de contraseñas, ni un escáner de seguridad, ni un sistema de exámenes antitrampa, ni un "
        "servicio de ejecución en hardware IBM. No entrena modelos de machine learning. No ataca sistemas reales: el "
        "oráculo de Grover conoce la respuesta porque nosotros se la damos.")
    g.callout("jurado", [
        "**«¿En una frase, qué aprende un estudiante con QubitGuard?»**",
        "Aprende que en cuántica no se manipulan probabilidades directamente sino amplitudes, que pueden ser negativas e "
        "interferir; y lo comprueba prediciendo y midiendo cómo Grover concentra la probabilidad en la respuesta, y cómo "
        "el ruido de un chip real la vuelve a dispersar.",
    ])


def chapter2(g: Guide, f: dict) -> None:
    g.h1("2. Computación cuántica mínima")
    g.p("Este capítulo da exactamente lo necesario para entender el resto. Nada más.")
    g.h2("2.1 Bit y qubit")
    g.p("Un **bit** clásico vale 0 o 1. Un **qubit** es un sistema físico (en los chips de IBM, un circuito "
        "superconductor enfriado casi al cero absoluto) cuyo estado se describe con dos números llamados "
        "**amplitudes**, uno para el resultado 0 y otro para el resultado 1:")
    g.code("|ψ⟩ = α·|0⟩ + β·|1⟩        con   α² + β² = 1   (en QubitGuard, α y β son reales)", size=9.5)
    g.p("La notación `|0⟩` (se lee «ket cero») solo es un nombre para el estado. En general las amplitudes son números "
        "complejos y la condición es `|α|² + |β|² = 1`; en todo lo que hace QubitGuard basta con números reales, "
        "**que pueden ser negativos**.")

    g.h2("2.2 Amplitud y probabilidad")
    g.p("Al medir un qubit sale 0 o 1. La probabilidad de cada resultado es **el cuadrado de su amplitud**:")
    g.code("P(0) = α²        P(1) = β²", size=9.5)
    g.table(["Amplitudes (α, β)", "P(0)", "P(1)", "Comentario"], [
        ["(1, 0)", "100 %", "0 %", "El qubit está en 0 con certeza."],
        ["(1/√2, 1/√2) ≈ (0,707; 0,707)", "50 %", "50 %", "Superposición uniforme: lo que deja una puerta H."],
        ["(1/√2, −1/√2)", "50 %", "50 %", "Mismas probabilidades, ¡pero otro estado! El signo importa al combinar."],
    ], [5.5, 1.8, 1.8, 6.9])
    g.callout("idea", [
        "**Probabilidad = amplitud².** El signo de la amplitud no cambia la probabilidad de ese resultado por sí solo, "
        "pero sí cambia lo que pasa cuando las amplitudes se suman o se restan. Ese es todo el truco de Grover.",
    ])

    g.h2("2.3 Superposición, sin mitos")
    g.p("Decimos que un qubit está en **superposición** cuando sus dos amplitudes son distintas de cero. Eso **no** "
        "significa que el qubit «sea 0 y 1 a la vez» en el sentido cotidiano, ni que «pruebe todas las respuestas a la "
        "vez». Significa que su estado se describe con amplitudes, y que al medir obtendremos **un solo** resultado, "
        "con probabilidad igual al cuadrado de la amplitud.")
    g.p("La típica analogía de la moneda girando sirve solo para la idea de «todavía no hay resultado», pero falla en "
        "lo importante: las probabilidades de una moneda nunca se cancelan entre sí. Las amplitudes sí. Una analogía "
        "más fiel es la de las ondas: dos ondas que se encuentran pueden reforzarse (crestas con crestas) o anularse "
        "(cresta con valle).")

    g.h2("2.4 Medición y shots")
    g.p("**Medir** convierte el estado cuántico en un resultado clásico (una cadena de bits) y destruye la "
        "superposición. Una sola medición da un solo resultado, así que para conocer la distribución hay que preparar "
        "y medir el mismo circuito muchas veces. Cada repetición se llama **shot** (disparo). Con 1024 shots obtenemos "
        "1024 resultados y los contamos: eso son los **counts** (conteos) que dibuja el histograma de la app.")

    g.h2("2.5 Interferencia")
    g.p("Las puertas cuánticas combinan amplitudes sumándolas y restándolas. Si dos contribuciones tienen el mismo "
        "signo, se refuerzan (**interferencia constructiva**); si tienen signo opuesto, se cancelan (**interferencia "
        "destructiva**). Un algoritmo cuántico útil está diseñado para que las respuestas incorrectas se cancelen y la "
        "correcta se refuerce. Grover lo consigue con dos reflexiones que se repiten (capítulo 4).")

    g.h2("2.6 Varios qubits")
    g.p("Con **n qubits** hay `N = 2^n` resultados posibles (estados de la base): con 2 qubits, `00, 01, 10, 11`; "
        "con 4 qubits, 16; con 6 qubits, 64. El estado completo es una lista de N amplitudes. Por eso simular "
        "exactamente un circuito cuesta memoria que crece como `2^n`: es fácil con 6 qubits y imposible con 300.")
    g.callout("cuidado", [
        "**Orden de los bits en Qiskit.** En una cadena como `\"10\"`, el bit de la **derecha** es el qubit 0 (`q_0`) y "
        "el de la izquierda es el qubit más alto. Por eso el oráculo de QubitGuard recorre `reversed(marked_state)`. "
        "Las claves de `counts` siguen la misma convención (bit más significativo a la izquierda).",
    ])

    g.h2("2.7 Las puertas que usamos")
    g.table(["Puerta", "Qué hace", "Para qué la usa QubitGuard"], [
        ["**H** (Hadamard)", "`|0⟩ → (|0⟩+|1⟩)/√2` y `|1⟩ → (|0⟩−|1⟩)/√2`. Crea o deshace superposición.", "Superposición uniforme inicial; parte del difusor; convertir una X controlada en una Z controlada."],
        ["**X** (NOT)", "Intercambia `|0⟩` y `|1⟩`.", "En el oráculo, convierte los ceros del estado marcado en unos, y luego los devuelve."],
        ["**Z**", "Deja `|0⟩` igual y cambia el signo de `|1⟩`.", "Es la idea de «cambiar un signo»: la base del oráculo de fase."],
        ["**CZ**", "Cambia el signo solo de `|11⟩`.", "Es la puerta nativa de 2 qubits del chip FakeFez; el transpilador lo traduce todo a ella."],
        ["**MCZ** (Z multicontrolada)", "Cambia el signo solo del estado `|11…1⟩`.", "Corazón del oráculo y del difusor. En el código se construye como `H` + `mcx` + `H` sobre el último qubit."],
        ["**MCX** (`mcx`)", "Invierte el qubit objetivo solo si todos los controles valen 1.", "Pieza con la que se arma la MCZ."],
        ["**Medición**", "Devuelve 0 o 1 y destruye la superposición.", "`measure_all()` al final del circuito."],
    ], [3.3, 6.2, 6.5], caption="Puertas cuánticas del proyecto.")
    g.p("Un **circuito cuántico** es una secuencia de puertas aplicadas a unos cables (uno por qubit) que se lee de "
        "izquierda a derecha, y que termina con mediciones. En el capítulo 3 verás el dibujo real del circuito de "
        "QubitGuard para 2 qubits.")

    g.h2("2.8 Mitos frecuentes")
    g.table(["Mito", "Realidad"], [
        ["«Una computadora cuántica prueba todas las respuestas a la vez.»", "Al medir sale un solo resultado. La ventaja viene de la interferencia, que hace probable el correcto. Grover necesita unas `π/4·√N` repeticiones, no una."],
        ["«Lo cuántico es exponencialmente más rápido para todo.»", "Depende del problema. Para búsqueda sin estructura, Grover es cuadrático. Shor (factorización) sí da una ventaja enorme, pero para otro problema."],
        ["«Ya se pueden romper contraseñas con computadoras cuánticas.»", "No existe hoy un ataque cuántico práctico a contraseñas ni a AES. Con corrección de errores, el coste real de Grover sería enorme."],
        ["«Más qubits siempre es mejor.»", "En hardware ruidoso (NISQ), un circuito más grande acumula más error. QubitGuard lo muestra: con 6 qubits y ruido, el éxito cae al nivel del azar."],
    ], [6, 10])
    g.callout("jurado", [
        "**«¿Qué es una amplitud y por qué no trabajan directamente con probabilidades?»**",
        "Una amplitud es el número que describe cada resultado posible; su cuadrado es la probabilidad. Puede ser negativa, "
        "y por eso dos caminos pueden cancelarse. Las probabilidades nunca se cancelan, así que con ellas no se podría "
        "explicar Grover: el oráculo solo cambia un signo, y eso es invisible si miras probabilidades.",
    ])


def chapter3(g: Guide, f: dict) -> None:
    v = f["versions"]
    g.h1("3. Qué es Qiskit y qué usamos de él")
    g.h2("3.1 Qiskit en dos párrafos")
    g.p("**Qiskit** es el kit de desarrollo (SDK) de código abierto de IBM para computación cuántica, en Python. "
        "Sirve para **construir** circuitos cuánticos, **optimizarlos** para un chip concreto (transpilar), "
        "**ejecutarlos** en simuladores o en procesadores cuánticos reales de IBM, y **leer** los resultados. "
        "Su código está en github.com/Qiskit/qiskit.")
    g.p(f"QubitGuard usa tres paquetes del ecosistema (versiones verificadas: qiskit {v[0]}, qiskit-aer {v[1]}, "
        f"qiskit-ibm-runtime {v[2]}; rangos en `requirements.txt`): `qiskit` (circuitos, transpilación, simulación "
        "exacta), `qiskit-aer` (simulador con shots y ruido) y `qiskit-ibm-runtime` (FakeFez y acceso a IBM Quantum).")
    g.h2("3.2 El flujo de trabajo")
    g.p("La documentación de IBM describe el trabajo con Qiskit en cuatro pasos (los «Qiskit patterns»). QubitGuard "
        "los sigue tal cual:")
    g.table(["Paso", "Qué significa", "En QubitGuard"], [
        ["1. Mapear (construir)", "Expresar el problema como un circuito.", "`build_grover_circuit` arma superposición, oráculo, difusor y medición."],
        ["2. Optimizar (transpilar)", "Adaptar el circuito al hardware o simulador de destino.", "`transpile(...)` en la web; `generate_preset_pass_manager(...)` en el CLI."],
        ["3. Ejecutar", "Correrlo con una primitiva o un simulador.", "`AerSimulator.run(...)` en la web; `SamplerV2(...).run(...)` en el CLI para IBM."],
        ["4. Post-procesar", "Convertir los resultados en una respuesta.", "Contar aciertos, calcular `measured_success`, dibujar histograma y explicación."],
    ], [3.6, 5.4, 7])

    g.h2("3.3 Mapa de piezas de Qiskit en el proyecto")
    g.table(["Pieza", "Paquete", "Qué hace", "Dónde se usa"], [
        ["`QuantumCircuit`", "qiskit", "Representa un circuito: qubits, puertas y mediciones.", "Oráculo, difusor y circuito completo (`experiments.py`)."],
        ["`compose`", "qiskit", "Pega un circuito dentro de otro.", "Repetir k veces oráculo + difusor."],
        ["`transpile`", "qiskit", "Reescribe el circuito con las puertas y conexiones del destino y lo optimiza.", "`_run_grover` (web)."],
        ["`generate_preset_pass_manager`", "qiskit.transpiler", "Crea una tubería de transpilación reutilizable de nivel 0 a 3.", "CLI `scripts/run_ibm_hardware.py`."],
        ["`Statevector`", "qiskit.quantum_info", "Calcula el estado exacto (todas las amplitudes) sin medir.", "Curva de éxito `grover_curve`."],
        ["`AerSimulator`", "qiskit_aer", "Simula la ejecución con shots y devuelve conteos; admite ruido.", "Laboratorio ideal y ruidoso."],
        ["`FakeFez`", "qiskit_ibm_runtime.fake_provider", "Backend falso: copia local de las características y el ruido de un chip IBM de 156 qubits.", "Modo con ruido (web) y modo `fake` (CLI)."],
        ["`QiskitRuntimeService`", "qiskit_ibm_runtime", "Conexión autenticada con IBM Quantum Platform.", "CLI, modos `dry-run` e `ibm`."],
        ["`SamplerV2`", "qiskit_ibm_runtime", "Primitiva que ejecuta circuitos en un procesador y devuelve muestras.", "CLI, modo `ibm`."],
    ], [3.9, 3.4, 5, 3.7], size=9, caption="Piezas de Qiskit usadas y su función.")

    g.h2("3.4 QuantumCircuit: construir el circuito")
    g.p("**Qué es:** el objeto central de Qiskit; guarda una lista ordenada de instrucciones sobre qubits. "
        "**Para qué:** describir el algoritmo. **Cómo se aplica aquí:** el oráculo, el difusor y el circuito completo "
        "son tres `QuantumCircuit`. Este es el código real de `app/quantum/experiments.py` (resumido solo en la "
        "validación de entradas):")
    g.code('''def _phase_oracle_for_basis_state(n_qubits, marked_state):
    oracle = QuantumCircuit(n_qubits, name="phase_oracle")
    zero_qubits = [i for i, bit in enumerate(reversed(marked_state)) if bit == "0"]
    for qubit in zero_qubits:
        oracle.x(qubit)                                   # ceros del objetivo -> unos
    oracle.h(n_qubits - 1)
    oracle.mcx(list(range(n_qubits - 1)), n_qubits - 1)   # H + MCX + H = Z multicontrolada
    oracle.h(n_qubits - 1)
    for qubit in zero_qubits:
        oracle.x(qubit)                                   # deshacer
    return oracle

def _diffuser(n_qubits):
    diffuser = QuantumCircuit(n_qubits, name="diffuser")
    diffuser.h(range(n_qubits))
    diffuser.compose(_phase_oracle_for_basis_state(n_qubits, "0" * n_qubits), inplace=True)
    diffuser.h(range(n_qubits))
    return diffuser

def build_grover_circuit(n_qubits, marked_state, iterations=None, *, measure=True):
    ...
    circuit = QuantumCircuit(n_qubits, name="grover_demo")
    circuit.h(range(n_qubits))                            # superposición uniforme
    oracle = _phase_oracle_for_basis_state(n_qubits, marked_state)
    diffuser = _diffuser(n_qubits)
    for _ in range(iterations):
        circuit.compose(oracle, inplace=True)
        circuit.compose(diffuser, inplace=True)
    if measure:
        circuit.measure_all()                             # barrera + registro clásico "meas"
    return circuit, iterations''', size=8, caption="app/quantum/experiments.py")
    g.bullets([
        "`h(range(n))` aplica Hadamard a todos los qubits; `x(q)` aplica NOT a un qubit.",
        "`mcx(controles, objetivo)` es la X multicontrolada. Rodeada por H en el objetivo se convierte en una Z "
        "multicontrolada, que cambia el signo de `|11…1⟩`.",
        "`compose(otro, inplace=True)` añade el circuito `otro` al final, sin crear una copia.",
        "`measure_all()` añade una barrera y mide todos los qubits en un registro clásico llamado `meas`. Ese nombre "
        "importa en hardware real: el CLI lee los resultados con `result[0].data.meas.get_counts()`.",
    ])
    g.p(f"Este es el dibujo real que produce `circuit.draw(\"text\")` para 2 qubits y estado marcado `11` "
        f"(k óptimo = {f['k_n2']}). Es el mismo texto que la app muestra en el campo `circuit_text`:")
    g.code(f["draw_n2"], size=7.5, caption="Salida real de Qiskit: build_grover_circuit(2, \"11\")")
    g.p("Se lee así: las dos H iniciales crean la superposición; el bloque `H · X(controlada) · H` sobre `q_1` es el "
        "oráculo (Z controlada sobre `11`); después viene el difusor (H, X, Z controlada, X, H); al final, la barrera "
        "`░` y las dos mediciones `M`.")

    g.h2("3.5 transpile y generate_preset_pass_manager: adaptar el circuito al chip")
    g.p("**Qué es transpilar.** Un chip real no entiende cualquier puerta. Solo ejecuta un conjunto pequeño de "
        f"**puertas nativas** y solo puede aplicar puertas de 2 qubits entre qubits **vecinos**. Por ejemplo, el "
        f"backend {f['fez_name']} (FakeFez) tiene {f['fez_qubits']} qubits y sus operaciones son "
        f"`{', '.join(op for op in f['fez_ops'] if op in ('cz', 'rz', 'sx', 'x', 'measure'))}` (más `delay`, `reset`, `id` y "
        "control de flujo). Una `mcx` no existe en el chip: hay que descomponerla en muchas puertas pequeñas.")
    g.p("Transpilar es reescribir el circuito para un destino concreto. El transpilador hace, en orden:")
    g.steps([
        "**Layout (colocación):** decide qué qubit físico del chip hace de cada qubit lógico.",
        "**Routing (enrutado):** si dos qubits que deben interactuar no son vecinos en el **mapa de conexiones**, inserta "
        "puertas SWAP para acercarlos. Cada SWAP cuesta 3 puertas de 2 qubits.",
        "**Traducción:** convierte cada puerta a las puertas nativas (por ejemplo, `mcx` y `h` en `cz`, `rz`, `sx`, `x`).",
        "**Optimización:** simplifica (cancela puertas que se anulan, fusiona rotaciones) para reducir la profundidad y el "
        "número de puertas de 2 qubits, que son las que más error introducen.",
    ])
    g.table(["Nivel", "Qué hace"], [
        ["0", "Sin optimización: solo lo necesario para que el circuito pueda ejecutarse."],
        ["1", "Optimización ligera. Es la que usa QubitGuard en el modo ideal."],
        ["2", "Optimización media, con más pasadas de simplificación."],
        ["3", "La más intensa: más pasadas y búsqueda de mejores colocaciones. Tarda más en compilar y suele dar circuitos más cortos. La usan el modo con ruido de la web y el CLI."],
    ], [2, 14], caption="Niveles de optimización del transpilador (optimization_level).")
    g.code('''# app/quantum/experiments.py, dentro de _run_grover
simulator = _noisy_backend() if noisy else AerSimulator(seed_simulator=seed)
# Aer ejecuta mcx de forma nativa; la base cx/rz/sx/x hace que depth y two_qubit_gates
# sean comparables con FakeFez.
target = {"backend": simulator} if noisy else {"basis_gates": ["cx", "rz", "sx", "x"]}
compiled = transpile(circuit, **target, seed_transpiler=seed,
                     optimization_level=3 if noisy else 1)

# scripts/run_ibm_hardware.py: misma idea, como tubería reutilizable
compiled = generate_preset_pass_manager(
    optimization_level=3, backend=backend, seed_transpiler=args.seed
).run(circuit)''', size=8, caption="Transpilación en la web y en el CLI")
    g.p("`generate_preset_pass_manager` crea un **pass manager**: una tubería con todas las pasadas del nivel "
        "elegido (layout, routing, traducción, optimización). Hace lo mismo que `transpile`, pero como objeto "
        "reutilizable; es la forma recomendada por IBM para preparar circuitos para Qiskit Runtime.")
    i4 = f["ideal4"]
    g.callout("ejemplo", [
        f"Con 2 qubits, el circuito transpilado a `cx/rz/sx/x` (nivel 1) tiene profundidad **{f['n2_depth']}** y "
        f"**{f['n2_2q']}** puertas de 2 qubits. Con 4 qubits y k = 3: profundidad **{i4['depth']}** y "
        f"**{i4['two_qubit_gates']}** puertas de 2 qubits.",
        "Con 3 qubits transpilado para FakeFez (nivel 3, seed 7, ejecutado con el CLI): profundidad **137**, **39** "
        "puertas de 2 qubits; colocación inicial en los qubits físicos `[24, 22, 23]` y final `[24, 23, 22]`. Que "
        "la colocación final sea distinta de la inicial indica que el enrutado permutó qubits por el camino.",
        "La **profundidad** es el número de capas de puertas que no pueden ejecutarse en paralelo: una medida del "
        "«largo» del circuito y, por tanto, del tiempo que los qubits deben resistir sin perder su estado.",
    ])

    g.h2("3.6 Statevector: simulación exacta, sin muestreo")
    g.p("**Qué es:** `qiskit.quantum_info.Statevector` calcula la lista completa de amplitudes de un circuito sin "
        "medir. **Para qué:** obtener probabilidades exactas, sin el «ruido» estadístico de tomar muestras. "
        "**Cómo se aplica:** la curva de éxito por iteración de la app (`POST /api/quantum/grover-curve`) construye "
        "el circuito sin medición para cada k y lee la probabilidad del estado marcado.")
    g.code('''for k in iterations:                     # k = 0 .. 2·óptimo
    circuit, _ = build_grover_circuit(n_qubits, marked_state, k, measure=False)
    probabilities.append(float(
        Statevector.from_instruction(circuit).probabilities()[int(marked_state, 2)]))''',
           size=8, caption="app/quantum/experiments.py, _grover_curve_cached")
    s_h, s_o, s_d = f["sv_steps"]
    fmtv = lambda vec: "[" + ", ".join(f"{x:+.2f}".replace("+0.00", " 0.00").replace("-0.00", " 0.00") for x in vec) + "]"
    g.p("Ejemplo real con 2 qubits y estado marcado `11`. Las amplitudes están en el orden `|00⟩, |01⟩, |10⟩, |11⟩`:")
    g.code(f"Tras las H:          {fmtv(s_h)}\n"
           f"Tras el oráculo:     {fmtv(s_o)}     <- solo cambia el signo de |11⟩\n"
           f"Tras el difusor:     {fmtv(s_d)}     <- toda la amplitud en |11⟩",
           size=8.5, caption="Salida real de Statevector, paso a paso")
    g.p("El resultado final es `−1` en `|11⟩`: el signo global no importa (la probabilidad es `(−1)² = 1`). El "
        "difusor de QubitGuard difiere del de los libros en esa fase global, que no cambia ninguna medición.")
    c = f["curve4"]["probability"]
    g.p("Y la curva real para 4 qubits, k = 0 a 6: " + "; ".join(f"k={k}: {fmt(p*100, 2)} %" for k, p in enumerate(c)) +
        ". La prueba `test_grover_curve_matches_formula_at_every_k` comprueba que, de 2 a 6 qubits, estos valores "
        "coinciden con la fórmula `sin²((2k+1)θ)` con un error máximo medido de `6,15·10⁻¹⁴`.")
    g.callout("cuidado", [
        "Statevector guarda `2^n` amplitudes. Con 6 qubits son 64 números, trivial. Con 50 qubits serían más de mil "
        "billones: por eso la simulación exacta solo sirve para circuitos pequeños, y por eso existen las computadoras cuánticas.",
    ])

    g.h2("3.7 AerSimulator: medir con shots")
    g.p("**Qué es:** `qiskit_aer.AerSimulator` es un simulador de alto rendimiento que **ejecuta** el circuito como "
        "si fuera un dispositivo: mide `shots` veces y devuelve conteos. **Para qué:** que el estudiante vea qué es "
        "un shot y que lo medido fluctúa alrededor de la teoría. **Cómo se aplica:** el laboratorio ideal usa "
        "`AerSimulator(seed_simulator=seed)` y lee `result().get_counts(compiled)`. La **semilla** (`seed`) hace "
        "que la misma petición dé los mismos conteos en el mismo entorno; la interfaz web no fija semilla, la API sí lo permite.")
    k0 = f["k0_4"]["counts"]
    g.callout("ejemplo", [
        f"**2 qubits, marcado `11`, 1024 shots:** conteos reales `{dict(f['n2_counts'])}`. Con N = 4 una sola "
        "iteración de Grover da el 100 % teórico.",
        f"**4 qubits, k = 0, 1024 shots (seed 7):** los 16 estados salen entre {min(k0.values())} y {max(k0.values())} "
        f"veces cada uno (lo esperado sería 64): es la distribución uniforme, una búsqueda a ciegas.",
        f"**4 qubits, marcado `1010`, k = 3, 2048 shots, seed 7:** el marcado salió {i4['counts']['1010']} veces, éxito "
        f"medido **{fmt(i4['measured_success'], 4)}** frente a **{fmt(i4['theoretical_success'], 4)}** teórico. La diferencia "
        "es azar de muestreo, no un error.",
    ])

    g.h2("3.8 FakeFez y AerSimulator.from_backend: el ruido de un chip real, en local")
    g.p("**Qué es un fake backend.** Es un objeto que vive en tu computadora y **imita** a un procesador real de IBM: "
        "tiene su número de qubits, sus puertas nativas, su mapa de conexiones y una **instantánea de datos de "
        "calibración** (tasas de error de cada puerta, tiempos de coherencia, errores de lectura). No se conecta a IBM "
        f"ni necesita cuenta. `FakeFez` imita al procesador `ibm_fez`, de {f['fez_qubits']} qubits.")
    g.p("**Para qué lo usamos.** `AerSimulator.from_backend(FakeFez())` crea un simulador que aplica el modelo de ruido "
        "de ese chip. Así el estudiante ve, sin acceso a hardware, que el ruido degrada el resultado. En el código se "
        "crea una sola vez y se reutiliza gracias a `lru_cache`:")
    g.code('''@lru_cache(maxsize=1)
def _noisy_backend() -> AerSimulator:
    return AerSimulator.from_backend(FakeFez())''', size=8.5, caption="app/quantum/experiments.py")
    g.callout("cuidado", [
        "FakeFez es una **simulación**. Sus datos son una foto antigua del chip, no la calibración de hoy. Di siempre "
        "«con el modelo de ruido FakeFez» o «simulación con ruido», nunca «en el chip de IBM».",
    ])

    g.h2("3.9 qiskit_ibm_runtime: el camino a un procesador real")
    g.p("Este paquete conecta con **IBM Quantum Platform** y ejecuta en procesadores reales. QubitGuard lo usa solo en "
        "el CLI `scripts/run_ibm_hardware.py`, nunca en el servidor web:")
    g.table(["Pieza", "Qué hace", "Uso en el CLI"], [
        ["`QiskitRuntimeService(...)`", "Abre la conexión autenticada con tu cuenta e instancia.", "`build_service()`: lee la cuenta guardada o las variables de entorno."],
        ["`service.backend(nombre)`", "Devuelve un procesador concreto.", "Si pasas `--backend`; el CLI comprueba que es real, operativo y con qubits suficientes."],
        ["`service.least_busy(...)`", "Elige el procesador con menos cola que cumpla filtros.", "`least_busy(operational=True, simulator=False, min_num_qubits=n)`."],
        ["`SamplerV2(mode=backend)`", "Primitiva que ejecuta circuitos y devuelve muestras.", "`.run([compiled], shots=shots)` envía un **job**."],
        ["`job.job_id()`, `job.result()`", "Identificador del trabajo y espera del resultado.", "Imprime el `job_id` al instante; luego lee `result()[0].data.meas.get_counts()`."],
        ["`service.job(id)`", "Recupera un job ya enviado.", "Modo `--job-id`: recupera sin reenviar."],
    ], [4.2, 5.6, 6.2], size=9)
    g.p("Un **job** es un trabajo enviado a la nube de IBM. Entra en una **cola** compartida con otros usuarios; cuando "
        "le toca, el procesador lo ejecuta y el resultado queda disponible con su `job_id`.")

    g.h2("3.10 Las primitivas: Sampler frente a Estimator")
    g.p("En Qiskit Runtime, las **primitivas** son las formas estándar de pedir algo a un procesador cuántico. Hay dos:")
    g.table(["Primitiva", "Qué devuelve", "Cuándo se usa", "¿En QubitGuard?"], [
        ["**Sampler**", "Muestras de las cadenas de bits medidas (conteos).", "Cuando la respuesta es un resultado de medición, como en Grover: queremos saber qué cadena sale.", "**Sí**: `SamplerV2` en el CLI."],
        ["**Estimator**", "Valores esperados de observables (por ejemplo, una energía).", "Química, optimización variacional (VQE), física: cuando se quiere un promedio de una magnitud.", "No."],
    ], [2.6, 4.4, 6, 3])
    g.p("La documentación más reciente de IBM presenta además primitivas nuevas; el CLI usa `SamplerV2`, que es la API "
        "instalada (Runtime 0.51) y la solicitada.")

    g.h2("3.11 Lo que NO usamos de Qiskit, y por qué")
    g.bullets([
        "**No usamos los bloques prefabricados de Grover de la biblioteca de circuitos.** El oráculo y el difusor se "
        "construyen a mano con `x`, `h` y `mcx` para que cada pieza se pueda explicar.",
        "**No usamos Estimator, ni QAE, ni algoritmos de machine learning cuántico.**",
        "**No configuramos mitigación de errores** en el CLI: el objetivo es ver el efecto del ruido, no ocultarlo.",
        "Lo que no necesita Qiskit no lo usa: Monte Carlo (NumPy), el explorador de amplitudes (fórmulas en el "
        "navegador) y el análisis de contraseñas (TypeScript local).",
    ])
    g.callout("jurado", [
        "**«Si la probabilidad tiene fórmula, ¿para qué usan Qiskit?»**",
        "Porque la fórmula describe un circuito ideal, y nosotros queremos que el estudiante vea un circuito de verdad: "
        "lo construimos, lo transpilamos y lo medimos con shots. Qiskit nos da lo que una fórmula no da: la profundidad "
        "y las puertas de 2 qubits tras transpilar, el muestreo con su variación estadística, el ruido de un chip "
        "simulado y el camino oficial para ejecutar el mismo circuito en un procesador de IBM.",
    ])


def chapter4(g: Guide, f: dict) -> None:
    g.h1("4. El algoritmo de Grover, paso a paso")
    g.h2("4.1 El problema de búsqueda")
    g.p("Tenemos N candidatos y uno solo es el correcto. Disponemos de un **oráculo**: una caja que, dado un candidato, "
        "responde «sí» o «no». No hay ninguna pista sobre dónde está (búsqueda **sin estructura**).")
    g.p("**Clásicamente**, lo mejor es probar candidatos en orden aleatorio sin repetir. La posición del secreto en ese "
        "orden es igual de probable en cualquier lugar de 1 a N, así que el número medio de intentos es el promedio de "
        "1, 2, …, N:")
    g.code("intentos esperados = (1 + 2 + … + N) / N = (N + 1) / 2", size=9.5)
    g.p("Con N = 16 son 8,5 intentos de media; con N = 1024, 512,5. Tras k intentos, la probabilidad de haberlo "
        "encontrado es `k/N`.")

    g.h2("4.2 La receta en tres bloques")
    g.steps([
        "**Superposición:** aplicar H a todos los qubits. Todos los estados quedan con amplitud `1/√N`.",
        "**Repetir k veces:** el **oráculo** cambia el signo del estado marcado y el **difusor** refleja todas las "
        "amplitudes respecto a su media. Cada repetición se llama **iteración de Grover**.",
        "**Medir.** Si k es el adecuado, sale el estado marcado con probabilidad alta.",
    ])

    g.h2("4.3 El oráculo de fase")
    g.p("El oráculo no «busca» nada: **cambia el signo de la amplitud del estado marcado** y deja el resto igual. "
        "Como `(−a)² = a²`, este paso por sí solo **no cambia ninguna probabilidad**. Prepara el terreno para el difusor.")
    g.p("Cómo lo hace el código para un estado marcado cualquiera, por ejemplo `10`:")
    g.steps([
        "Aplica X a los qubits cuyo bit en el objetivo es 0 (leyendo la cadena al revés por el orden de Qiskit). "
        "Así el estado marcado se convierte temporalmente en `11`.",
        "Aplica una Z multicontrolada (H + `mcx` + H en el último qubit), que cambia el signo solo de `11…1`.",
        "Deshace las X del paso 1.",
    ])
    g.code(f["oracle_10"], size=8, caption="Salida real de Qiskit: _phase_oracle_for_basis_state(2, \"10\")")
    g.p("Para el marcado `10`, el qubit `q_0` vale 0, así que recibe la X antes y después de la Z controlada.")
    g.callout("cuidado", [
        "El oráculo **conoce la respuesta**: se la damos al construir el circuito. En un ataque real habría que "
        "programar como circuito reversible la función que verifica una clave (un hash o un cifrado), con muchísimos "
        "qubits y puertas. QubitGuard ilustra la mecánica de Grover; no descubre secretos.",
    ])

    g.h2("4.4 El difusor: inversión sobre la media")
    g.p("El difusor aplica H a todos los qubits, después el oráculo del estado `00…0` y después H otra vez. Su efecto "
        "es una **reflexión respecto a la media**: cada amplitud `a` se convierte en `2·media − a` (salvo un signo global "
        "que no afecta a las mediciones).")
    g.p("Intuición: después del oráculo, el marcado tiene amplitud negativa y queda muy por debajo de la media. Al "
        "reflejar respecto a la media, lo que estaba muy por debajo pasa a estar muy por encima.")
    g.callout("ejemplo", [
        "**N = 16 (4 qubits), una iteración.** Al empezar, las 16 amplitudes valen `1/4 = 0,25`.",
        "- Oráculo: el marcado pasa a `−0,25`; los otros 15 siguen en `0,25`.",
        "- Media: `(15 · 0,25 − 0,25) / 16 = 0,21875`.",
        "- Difusor: los otros pasan a `2 · 0,21875 − 0,25 = 0,1875`; el marcado, a `2 · 0,21875 + 0,25 = 0,6875`.",
        "- Probabilidad del marcado: `0,6875² ≈ 0,4727`, es decir, **47,3 %** (antes era 6,25 %). Coincide con el valor "
        "de Statevector para k = 1.",
    ])

    g.h2("4.5 La geometría: Grover es una rotación")
    g.p("Todo el algoritmo ocurre en un plano con dos ejes: `|ω⟩`, el estado marcado, y `|s'⟩`, la superposición "
        "uniforme de todos los demás. El estado inicial forma un ángulo pequeño θ con `|s'⟩`, donde:")
    g.code("sin θ = 1/√N        es decir        θ = asin(1/√N)", size=9.5)
    g.p("El oráculo es una reflexión respecto a `|s'⟩` y el difusor es una reflexión respecto al estado inicial. "
        "**Dos reflexiones seguidas equivalen a una rotación** del doble del ángulo entre los ejes: cada iteración "
        "gira el estado **2θ** hacia `|ω⟩`. Después de k iteraciones el ángulo es `(2k+1)θ` y la probabilidad de medir "
        "el marcado es su proyección al cuadrado:")
    g.code("P(k) = sin²((2k + 1) · θ)", size=10)
    g.figure(g.fig["geometry"], 10, "Cada iteración gira el estado 2θ. Con N = 16, θ ≈ 14,5°; tras 3 iteraciones el "
             "ángulo es 101,3°, cerca de 90°, y la probabilidad es 96,1 %. Figura generada con las fórmulas del proyecto.")

    g.h2("4.6 El número óptimo de iteraciones")
    g.p("Queremos que el ángulo `(2k+1)θ` quede lo más cerca posible de 90° (π/2), donde `sin² = 1`. Despejando, "
        "`k ≈ π/(4θ) − 1/2`. QubitGuard usa `optimal_iterations(n) = floor(π/(4θ))`; para n de 2 a 6 da lo mismo que "
        "redondear `π/(4θ) − 1/2`. Para N grande, θ ≈ 1/√N y por eso `k ≈ (π/4)·√N`.")
    rows = []
    for n, N, th, ko, p in f["opt_table"]:
        rows.append([str(n), str(N), f"{fmt(th, 2)}°", str(ko), f"{fmt(p*100, 2)} %", fmt((N + 1) / 2, 1),
                     fmt(math.pi / 4 * math.sqrt(N), 2)])
    g.table(["n", "N = 2^n", "θ", "k óptimo", "P(k óptimo)", "Clásico (N+1)/2", "π/4·√N"], rows,
            [1, 1.8, 2, 2, 2.6, 3, 2.2], caption="Iteraciones óptimas calculadas con optimal_iterations y theoretical_success.")
    g.callout("idea", [
        "Memoriza la fila de 4 qubits: **N = 16, k óptimo = 3, probabilidad 96,1 %**, frente a 8,5 intentos clásicos de "
        "media. Es la demo recomendada para el jurado.",
    ])

    g.h2("4.7 Por qué pasarse baja la probabilidad")
    g.p("Como Grover es una rotación, si seguimos iterando después del óptimo la flecha **se pasa** de `|ω⟩` y sigue "
        "girando: la probabilidad baja. Con 4 qubits, k = 4 da 58,2 % y k = 6, apenas 2,0 %. En la búsqueda clásica, en "
        "cambio, más intentos nunca empeoran el resultado. Esta es una de las ideas más contraintuitivas que enseña la app.")
    g.figure(g.fig["curve"], 14, "Probabilidad de éxito por iteración con 4 qubits, calculada con Statevector "
             "(grover_curve(4)), frente a una búsqueda al azar con las mismas k consultas.")

    g.h2("4.8 Ventaja cuadrática, no exponencial")
    rows = []
    for row in f["compare40"][:3] + [f["compare128"][-1]]:
        b = row["bits"]
        if b <= 52:
            rows.append([str(b), f"{row['N']:,}".replace(",", "."), fmt(row["classical_expected"], 1),
                         str(row["grover_iterations"]), fmt(row["speedup"], 1)])
        else:
            rows.append([str(b), f"{row['N']:.2e}".replace(".", ","), f"{row['classical_expected']:.2e}".replace(".", ","),
                         f"{row['grover_iterations']:.2e}".replace(".", ","), f"{row['speedup']:.2e}".replace(".", ",")])
    g.table(["Bits", "N", "Intentos clásicos (N+1)/2", "Iteraciones de Grover", "Razón"], rows, [1.4, 3.6, 4, 3.8, 3.2],
            caption="Salida real de compare(): razón entre consultas, no entre tiempos.")
    g.p("Grover reduce `N` a `√N`: es una **ventaja cuadrática**. Es grande, pero no hace fácil un problema enorme: "
        "con 128 bits siguen haciendo falta unas `1,45·10^19` iteraciones, cada una con un oráculo costoso y en hardware "
        "tolerante a fallos que hoy no existe. Por eso **duplicar la longitud de la clave restaura el margen**: AES-256 "
        "conserva unos 128 bits de seguridad idealizada frente a Grover. El algoritmo que sí rompe RSA y curvas "
        "elípticas es el de **Shor**, que QubitGuard explica en sus tarjetas pero no simula.")
    g.callout("cuidado", [
        "La columna «razón» compara **consultas al oráculo**, no segundos. Una iteración cuántica tiene su propio coste, "
        "y simular un circuito consume CPU clásica.",
    ])

    g.h2("4.9 Ejemplo resuelto a mano: 2 qubits")
    g.p("Estado marcado `11`, N = 4. Amplitudes en el orden `|00⟩, |01⟩, |10⟩, |11⟩`:")
    g.table(["Paso", "|00⟩", "|01⟩", "|10⟩", "|11⟩", "Qué pasó"], [
        ["Inicio", "1", "0", "0", "0", "Todos los qubits en 0."],
        ["Tras H", "0,5", "0,5", "0,5", "0,5", "Superposición uniforme: cada estado con 25 %."],
        ["Tras oráculo", "0,5", "0,5", "0,5", "−0,5", "Cambia el signo del marcado. Probabilidades: aún 25 % cada uno."],
        ["Media", "", "", "", "", "(0,5 + 0,5 + 0,5 − 0,5) / 4 = 0,25"],
        ["Tras difusor", "0", "0", "0", "1", "2·0,25 − 0,5 = 0 para los otros; 2·0,25 + 0,5 = 1 para el marcado."],
    ], [2.6, 1.3, 1.3, 1.3, 1.3, 8.2])
    g.p("Con la fórmula: θ = asin(1/√4) = asin(0,5) = 30°. Con k = 1, `(2·1+1)·30° = 90°` y `sin²(90°) = 1`. "
        f"Probabilidad 100 %. Qiskit confirma: Statevector final `{fmtv_short(f['sv_steps'][2])}` (signo global "
        f"irrelevante) y AerSimulator con 1024 shots da `{dict(f['n2_counts'])}`.")
    g.callout("jurado", [
        "**«¿Por qué no iteran más veces para ir a lo seguro?»**",
        "Porque Grover es una rotación: después del óptimo la probabilidad baja. Con 4 qubits, 3 iteraciones dan 96,1 %, "
        "pero 6 iteraciones dan 2,0 %, menos que adivinar con una sola consulta. Además, en hardware real cada iteración "
        "añade puertas y ruido.",
    ])


def fmtv_short(vec) -> str:
    return "[" + ", ".join("0" if abs(x) < 1e-9 else f"{x:g}" for x in vec) + "]"


def chapter5(g: Guide, f: dict) -> None:
    g.h1("5. Cómo funciona el software por dentro")
    g.h2("5.1 Arquitectura")
    g.p("QubitGuard es **un solo proceso**: un servidor FastAPI (Python) que sirve a la vez la API y la página web. El "
        "navegador ejecuta la interfaz en TypeScript compilado a JavaScript, sin frameworks y sin dependencias de "
        "ejecución. Los experimentos cuánticos corren en el servidor con Qiskit y Aer.")
    g.code('''                      mismo origen y puerto (por defecto :8000)
 ┌───────────────────┐  HTTP/JSON  ┌──────────────────────┐  llamadas  ┌────────────────────────┐
 │ Navegador         │ ──────────▶ │ FastAPI (un proceso) │ ─────────▶ │ Qiskit / Aer / FakeFez │
 │ TypeScript + SVG  │ ◀────────── │ rutas, validación,   │ ◀───────── │ NumPy (Monte Carlo)    │
 │ contraseñas y     │             │ cabeceras, caché,    │  locales   │ app/quantum/           │
 │ progreso locales  │             │ archivos estáticos   │            │   experiments.py       │
 └───────────────────┘             └──────────────────────┘            └────────────────────────┘

 Aparte, fuera del servidor web (requiere cuenta IBM):
 scripts/run_ibm_hardware.py -> mismo build_grover_circuit -> generate_preset_pass_manager
                                                          -> SamplerV2 -> IBM Quantum Platform''',
           size=7.2, caption="Arquitectura (adaptada del README)")
    g.callout("idea", [
        "Tres capas: **navegador** (interfaz, explorador, contraseñas y progreso, todo local), **FastAPI** (valida y "
        "enruta) y **motor cuántico** (Qiskit, Aer, FakeFez y NumPy). El CLI de hardware es un programa aparte: "
        "arrancar la web no crea cuentas ni gasta cuota de IBM.",
    ])
    g.h3("Recorrido de una ejecución")
    g.steps([
        "El usuario hace su predicción y pulsa **Ejecutar circuito**. `initLab` (en `main.ts`) valida los campos e impide "
        "envíos dobles.",
        "`api.grover(...)` envía `POST /api/quantum/grover` con n, estado marcado, iteraciones, shots, semilla y ruido. "
        "**Nunca viaja una contraseña.**",
        "`app/api/quantum.py` valida el cuerpo con Pydantic (`extra=\"forbid\"`: cualquier campo extra da error 422).",
        "`experiments.run_grover` vuelve a validar, construye el circuito, lo transpila, lo simula y mide profundidad y "
        "puertas de 2 qubits.",
        "Vuelve un JSON con conteos, éxito medido y teórico, métricas y el dibujo del circuito.",
        "`renderResult` dibuja histograma, curva y circuito; `observe` compara con la predicción; `afterRun` da XP, "
        "marca pasos de la ruta y logros.",
    ])

    g.h2("5.2 Los endpoints de la API")
    g.table(["Método y ruta", "Entrada", "Devuelve"], [
        ["GET `/api/health`", "Nada.", "`{\"status\": \"ok\", ...}`"],
        ["GET `/api/quantum/capabilities`", "Nada.", "Versiones de Qiskit y Aer, límites (2–6 qubits, 8192 shots), modelo de ruido y `physical_hardware: false`."],
        ["POST `/api/quantum/grover`", "`n_qubits` 2–6, `marked_state` (n bits), `iterations` 0–20 o null (óptimo), `shots` 1–8192, `seed` o null, `noisy`.", "`counts`, `measured_success`, `theoretical_success`, `optimal_iterations`, `depth`, `two_qubit_gates`, `backend`, `circuit_text`."],
        ["POST `/api/quantum/grover-curve`", "`n_qubits` 2–6.", "Probabilidad exacta (Statevector) para k = 0 … 2·óptimo."],
        ["POST `/api/quantum/montecarlo`", "`n_bits` 2–20, `trials` 1–5000, `seed`.", "Media medida, media esperada `(N+1)/2` e histograma."],
        ["GET `/api/quantum/compare?max_bits=40`", "`max_bits` 2–256.", "Tabla teórica clásico frente a Grover, cada 4 bits más 128/256 y el pedido."],
        ["GET `/api/learn/pqc`", "Nada.", "Tarjetas de criptografía poscuántica (desde `app/content/pqc.json`)."],
        ["GET `/api/learn/limitations`", "Nada.", "Limitaciones que ve la interfaz (desde `limitations.json`)."],
        ["GET `/api/challenge/daily?date=AAAA-MM-DD`", "Fecha (UTC si se omite).", "Pregunta del día, 4 opciones, índice correcto y explicación."],
    ], [5, 5.2, 5.8], size=8.8, caption="Contrato de la API (docs/API.md).")
    g.p("Errores: **422** si la entrada es inválida (rango, longitud del estado, campo extra); **503** con "
        "`Retry-After: 2` si el simulador con ruido está ocupado; **500** genérico, sin detalles internos.")

    g.h2("5.3 Módulos del backend y del frontend")
    g.table(["Archivo", "Rol"], [
        ["`app/main.py`", "Crea la app FastAPI, monta los routers, gzip, cabeceras de seguridad, política de caché y lista blanca de archivos públicos; traduce `SimulatorBusy` a 503."],
        ["`app/api/quantum.py`", "Modelos Pydantic (`GroverRequest`, `CurveRequest`, `MonteCarloRequest`) y rutas cuánticas; convierte `ValueError` en 422."],
        ["`app/quantum/experiments.py`", "Todo lo cuántico y estadístico: `build_grover_circuit`, `run_grover`, `grover_curve`, `montecarlo`, `compare`, `capabilities`."],
        ["`app/api/challenge.py`", "Reto diario determinista; no importa Qiskit."],
        ["`app/api/learn.py` y `app/content/*.json`", "Contenido educativo, leído una vez al arrancar."],
        ["`scripts/run_ibm_hardware.py`", "CLI para FakeFez o IBM Quantum (capítulo 6)."],
    ], [5.2, 10.8], size=9, caption="Backend (Python).")
    g.table(["Módulo", "Rol"], [
        ["`main.ts`", "Flujo de la interfaz: laboratorio, Predice-Observa-Explica, reto, Monte Carlo, comparación y navegación."],
        ["`api.ts`", "Cliente HTTP del contrato; caché por sesión con deduplicación de peticiones simultáneas; respeta `Retry-After`."],
        ["`charts.ts`", "Gráficos SVG propios (`barChart`, `lineChart`, `mount`) con título, descripción y tabla equivalente."],
        ["`explorer.ts`", "Explorador de amplitudes en vivo (fórmulas locales, sin servidor)."],
        ["`progress.ts`", "Ruta, logros, niveles, XP y racha en `localStorage`."],
        ["`password.ts`", "Analizador de fortaleza (funciones puras)."],
        ["`generator.ts` y `wordlist.ts`", "Generador de contraseñas con `crypto.getRandomValues`; lista de 2072 palabras."],
        ["`password-ui.ts`", "Conecta analizador y generador con la página; meta educativa de 128 bits."],
        ["`ambience.ts`, `dom.ts`", "Fondo ASCII, intro con el logo IBM en ASCII y utilidades de animación."],
    ], [4.2, 11.8], size=9, caption="Frontend (TypeScript en frontend/src/).")

    g.h2("5.4 Predice, Observa, Explica (POE)")
    g.p("Es la metodología central. El estudiante **se compromete** con una respuesta antes de ver el resultado, porque "
        "comparar una predicción propia con la realidad enseña más que leer la respuesta.")
    g.steps([
        "**Predice:** elige estimar la *probabilidad de acertar* (deslizador) o el *número de iteraciones óptimas*. La "
        "referencia visible es el azar, `1/N`.",
        "**Observa:** aparecen tres cifras en una misma escala: tu predicción, la teoría `sin²((2k+1)θ)` y lo medido. "
        "Debajo, histograma, curva, métricas del circuito transpilado y el circuito en texto.",
        "**Explica:** una lista corta interpreta el resultado: dónde quedó k respecto al óptimo, qué lograría el azar "
        "con las mismas consultas (`k/N`) y, con ruido, por qué lo medido cae respecto a la teoría.",
    ])
    g.table(["Modo de predicción", "Exacta", "Cerca", "Encaminada"], [
        ["Probabilidad (sin ruido se compara con la teoría; con ruido, con lo medido)", "≤ 3 puntos: 25 XP", "≤ 10 puntos: 15 XP", "≤ 25 puntos: 5 XP"],
        ["Iteraciones óptimas", "Exacto: 25 XP", "A 1 de distancia: 10 XP", "—"],
    ], [7, 3, 3, 3], size=9, caption="Puntuación de las predicciones (main.ts, observe). Solo puntúan 5 predicciones al día.")

    g.h2("5.5 Explorador de amplitudes")
    g.p("Sección «Amplitudes». El estudiante mueve **qubits** (2 a 10) e **iteraciones** (hasta `2·óptimo + 2`) y ve en "
        "tiempo real las barras de amplitud del marcado frente al resto con la línea de la media, el vector de estado "
        "girando en el plano `|ω⟩`/`|s'⟩`, y la comparación Grover frente al azar. El botón **Siguiente iteración** "
        "anima la iteración en dos tiempos: primero el oráculo (el ángulo pasa de φ a −φ) y luego el difusor (de −φ a "
        "φ + 2θ).")
    g.p("**No usa Qiskit ni el servidor:** calcula con fórmulas exactas en el navegador, para responder al instante:")
    g.code('''θ = asin(1/√N)          φ = (2k+1)·θ
amplitud del marcado        = sin φ
amplitud de cada uno de los N−1 restantes = cos φ / √(N−1)
media = (marcado + (N−1)·resto) / N''', size=8.5, caption="frontend/src/explorer.ts (amplitudesAt)")
    g.p("Estas fórmulas salen de escribir el estado como `sin φ·|ω⟩ + cos φ·|s'⟩`. Por eso puede llegar a 10 qubits "
        "aunque el laboratorio con circuitos llegue a 6. Coinciden con la curva de Statevector de 2 a 6 qubits.")
    if (CAPTURAS / "07_explorador.png").exists():
        g.figure(CAPTURAS / "07_explorador.png", 12.5, "Explorador de amplitudes con 4 qubits y 2 iteraciones: "
                 "90,8 % con Grover frente a 12,5 % al azar (captura real).")

    g.h2("5.6 Monte Carlo con NumPy")
    mc = f["mc10"]
    g.p("La sección «Azar frente a Grover» simula la búsqueda clásica. Ordenar los N candidatos al azar y contar en qué "
        "posición aparece el secreto es lo mismo que elegir esa posición **uniformemente entre 1 y N**: en una "
        "permutación aleatoria, el secreto tiene la misma probabilidad de estar en cualquier lugar. Así el código "
        "evita construir permutaciones de hasta un millón de elementos (2^20):")
    g.code('''attempts = np.random.default_rng(seed).integers(1, N + 1, size=trials)''', size=8.5,
           caption="app/quantum/experiments.py, montecarlo")
    g.p(f"Después agrupa los intentos en un histograma (un grupo por intento hasta 64 candidatos; 32 grupos para más) y "
        f"devuelve la media frente a `(N+1)/2`. Ejemplo real: 10 bits, 5000 ensayos, seed 1 → media "
        f"**{fmt(mc['mean_attempts'], 3)}** frente a **{fmt(mc['expected_attempts'], 1)}** esperados. El generador de "
        "NumPy es reproducible con semilla y **no es criptográfico**: sirve para estadística, no para secretos.")

    g.h2("5.7 Comparación clásico frente a Grover")
    g.p("`compare(max_bits)` calcula con fórmulas, hasta 256 bits, `(N+1)/2` frente a `π/4·√N` y su razón. No simula "
        "circuitos de ese tamaño. Por encima de 52 bits usa números de coma flotante para mostrar magnitudes.")

    g.h2("5.8 Reto diario determinista")
    g.p("Cada día hay una pregunta nueva de tres tipos: iteraciones óptimas, probabilidad de éxito o comparación "
        "clásico/Grover. No hay base de datos: la fecha ISO (por ejemplo `2026-10-24`) pasa por **SHA-256**, los primeros "
        "8 bytes del resumen son la semilla de `random.Random`, y de ahí salen el tipo, los parámetros y el orden de "
        "las 4 opciones. **La misma fecha produce siempre el mismo reto** para todos.")
    g.p("La interfaz pide el reto de la **fecha local** del navegador (la API usa UTC si no recibe fecha). La respuesta "
        "correcta viaja en el JSON a propósito: es un juego para aprender, no un examen.")

    g.h2("5.9 Racha, XP, niveles y logros")
    g.bullets([
        "Todo se guarda en el `localStorage` del navegador, clave `qubitguard.progress.v1`. No hay cuentas ni servidor "
        "de progreso; es editable y no sirve para rankings.",
        "**Reto diario:** acertar da 50 XP, intentarlo 15 XP, más un extra de racha de 10 XP por día seguido (máximo 50). "
        "Solo puntúa la primera respuesta de cada fecha; se guardan las últimas 60 fechas.",
        "**Racha:** días locales consecutivos con el reto resuelto; se rompe si no se respondió ni hoy ni ayer.",
        "**Niveles (7):** el nivel n requiere `60·(n−1)²` XP acumulados, de «Aprendiz de qubits» a «Guardián cuántico».",
        "**Actividades:** primera simulación de Grover del día (15 XP), primera con ruido (10 XP), primera Monte Carlo (10 XP).",
        "**Ruta (5 pasos):** superposición, oráculo, amplificación, ruido y criptografía. **Logros (6):** primera "
        "predicción, primera predicción exacta, ver el ruido borrar a Grover, contraseña de 128+ bits, tres días seguidos "
        "y ruta completa.",
    ])

    g.h2("5.10 Generador de contraseñas")
    g.p("Corre **solo en el navegador**: no hay ningún endpoint que reciba contraseñas, y una prueba automática lo "
        "comprueba. Tiene tres modos: caracteres (8 a 128, clases elegibles), frase de palabras (3 a 12 palabras de una "
        "lista de 2072) y «mejorar la mía» (inserta caracteres aleatorios en una base).")
    g.h3("De dónde sale el azar")
    g.code('''export function randomInt(maxExclusive: number): number {
  const limit = UINT32_RANGE - (UINT32_RANGE % maxExclusive);
  const sample = new Uint32Array(1);
  do {
    globalThis.crypto.getRandomValues(sample);   // CSPRNG del sistema operativo
  } while (sample[0]! >= limit);                 // muestreo por rechazo: sin sesgo de módulo
  return sample[0]! % maxExclusive;
}''', size=8.5, caption="frontend/src/generator.ts (sin la validación de entrada)")
    g.bullets([
        "`crypto.getRandomValues` es el **generador criptográficamente seguro (CSPRNG)** del navegador, alimentado por "
        "la entropía del sistema operativo. No admite semilla.",
        "**Sesgo de módulo:** si tomas un número de 0 a 2^32 − 1 y calculas el resto entre 94, los primeros valores salen "
        "un poco más a menudo porque 2^32 no es múltiplo de 94. El código descarta la «cola» que sobra y vuelve a "
        "sortear, así cada carácter es exactamente equiprobable.",
        "Las pruebas prohíben `Math.random`, red y almacenamiento en estos módulos.",
    ])
    g.h3("Por qué conocer el código no permite reproducir una contraseña")
    g.p("El código es público, y eso está bien: en criptografía se asume que el atacante conoce el algoritmo "
        "(principio de Kerckhoffs). La seguridad no está en ocultar el código sino en que **los bits aleatorios son "
        "impredecibles**: salen del sistema operativo en el momento de generar, no hay semilla que adivinar, no se "
        "guardan y no se envían. Hacer ingeniería inversa del generador solo revela *cómo* se combinan esos bits, no "
        "*cuáles* salieron. Es la diferencia con Monte Carlo, donde la semilla de NumPy hace los resultados "
        "reproducibles a propósito.")
    g.h3("Entropía")
    g.p("La **entropía** mide cuántas posibilidades igual de probables hay, en bits: `bits = log2(número de "
        "posibilidades)`. Cada bit duplica el trabajo de un atacante.")
    g.table(["Caso", "Cálculo", "Bits"], [
        ["16 caracteres de los 94 ASCII imprimibles", "`16 · log2(94) = 16 · 6,555`", "≈ 104,9 (algo menos al exigir cada clase)"],
        ["Una palabra de la lista", "`log2(2072)`", "≈ 11,02"],
        ["Frase de 12 palabras", "`12 · 11,02`", "≈ 132,2"],
        ["Mejorar la mía: cada carácter insertado", "`log2(94)`", "≈ 6,55 (la base cuenta como conocida: 0 bits)"],
    ], [5.6, 5.2, 5.2], size=9)
    g.p("Con clases obligatorias (por ejemplo, exigir al menos una mayúscula), no todas las cadenas valen; el código "
        "cuenta las válidas con **inclusión-exclusión** en escala logarítmica. La meta educativa es **128 bits**; la "
        "interfaz muestra la mitad (`groverBits = bits / 2`) como bits efectivos frente a un Grover idealizado, porque "
        "Grover convierte `2^b` intentos en unos `2^(b/2)`. Alcanzar la meta no certifica nada: es una referencia didáctica.")

    g.h2("5.11 Analizador de contraseñas")
    g.p("`analyze` (en `password.ts`) estima `longitud · log2(alfabeto detectado)` (26 minúsculas, 26 mayúsculas, 10 "
        "dígitos, 33 símbolos, ~100 otros) y **resta penalizaciones** por palabras comunes (también con «leet» como "
        "`p4ssw0rd`), años, recorridos de teclado, secuencias y repeticiones, sin contar dos veces el mismo carácter. "
        "Clasifica: menos de 28 bits muy débil, menos de 40 débil, menos de 64 razonable, menos de 128 fuerte y 128 o más "
        "muy fuerte.")
    g.p("El tiempo clásico supone una tasa de intentos elegida (por defecto 10^10 por segundo) y la mitad del espacio. "
        "El tiempo cuántico supone **un millón de iteraciones de Grover por segundo**, un escenario hipotético y muy "
        "optimista, no una medición. Es una heurística: no consulta filtraciones ni diccionarios completos.")

    g.h2("5.12 Seguridad")
    g.table(["Medida", "Detalle"], [
        ["Lista blanca de archivos", "Solo se publican `/`, `/index.html`, `/styles.css`, `/dist/`, `/assets/`, la API y su documentación. `/src/`, `/node_modules/` y `/package.json` dan 404. La ruta se normaliza antes de comprobarla (evita `..`)."],
        ["CSP", "`default-src 'self'`: sin scripts de terceros; se permiten estilos inline e imágenes `data:`. `/docs`, `/redoc` y `/openapi.json` quedan fuera de esta CSP."],
        ["Otras cabeceras", "`X-Content-Type-Options: nosniff`, `Referrer-Policy: no-referrer`, `X-Frame-Options: DENY`, `Cross-Origin-Opener-Policy`, `Permissions-Policy`. HSTS no se envía en local (HTTP); se activa en el borde al publicar por HTTPS."],
        ["Caché HTTP", "Simulaciones `no-store`; `learn`, `compare` y `capabilities` `max-age=3600`; reto diario `max-age=300`; estáticos `no-cache` con ETag."],
        ["Validación", "Pydantic con rangos y `extra=\"forbid\"`; el motor vuelve a validar. La semilla se limita a 2^32 − 1 para evitar un desbordamiento que daba error 500."],
        ["Límite de simulaciones con ruido", "Semáforo de 2 simulaciones FakeFez simultáneas por proceso; espera hasta 10 s y, si no hay cupo, 503 con `Retry-After: 2`."],
        ["Errores", "Un 500 sale con mensaje genérico y con las cabeceras de seguridad."],
        ["Lo que no hay", "Autenticación ni límite de peticiones por IP: para una publicación estable se configura en el borde (Cloudflare)."],
    ], [4, 12], size=9, caption="Medidas de seguridad (app/main.py y README, sección 8).")
    g.callout("jurado", [
        "**«¿La contraseña que escribo se envía a su servidor?»**",
        "No. El analizador y el generador corren en el navegador. La API no tiene ningún campo de contraseña, los "
        "cuerpos rechazan campos extra y hay pruebas que verifican que ningún endpoint acepta contraseñas y que esos "
        "módulos no usan red ni almacenamiento. Lo único que sale de tu navegador si lo haces tú es copiarla al portapapeles.",
    ])


def chapter6(g: Guide, f: dict) -> None:
    g.h1("6. Cómo se ejecutaría en una computadora cuántica real de IBM")
    g.callout("cuidado", [
        "**No se ejecutó nada en hardware real.** No se envió ningún job a IBM: no había cuenta configurada. Lo medido "
        "es FakeFez en local. Los modos `dry-run` e `ibm` del CLI se probaron con dobles locales, sin red. Este capítulo "
        "explica cómo se haría, con los comandos reales.",
    ], title="estado actual")
    g.h2("6.1 Qué cambia entre un simulador y un procesador real")
    g.table(["", "Simulador (lo que hicimos)", "Procesador IBM (pendiente)"], [
        ["Dónde corre", "En la CPU de nuestra computadora.", "En un chip superconductor en la nube de IBM."],
        ["Ruido", "Ninguno (ideal) o un modelo (FakeFez).", "Ruido real del día, según la calibración vigente."],
        ["Acceso", "Inmediato, sin cuenta.", "Cuenta, instancia, cola compartida y cuota de tiempo."],
        ["Circuito", "El mismo `build_grover_circuit`.", "El mismo `build_grover_circuit`, transpilado para ese chip."],
    ], [3, 6.5, 6.5], size=9.5)
    g.h2("6.2 Cuenta en IBM Quantum Platform")
    g.steps([
        "Crear un IBMid e iniciar sesión en IBM Quantum Platform (quantum.cloud.ibm.com).",
        "En **Instances**, crear o elegir una instancia del **plan Open** (gratuito; la guía del proyecto usa la región "
        "`us-east`) y copiar su identificador **CRN**.",
        "Crear una **API key** en el panel de la plataforma.",
    ])
    g.p("La documentación oficial consultada por el equipo indica que el plan Open da hasta **10 minutos de tiempo QPU "
        "por ventana móvil de 28 días**. Hay que verificarlo en el panel antes de usarlo. El tiempo en cola no cuenta "
        "como tiempo QPU.")

    g.h2("6.3 Credenciales seguras")
    g.p("Regla de oro: **la API key nunca va en el código, ni en un argumento de la línea de comandos, ni en el "
        "repositorio.** El CLI rechaza cualquier argumento desconocido con el mensaje «No se aceptan tokens por CLI», no "
        "guarda cuentas y no escribe credenciales en sus JSON ni en sus mensajes. Hay dos formas de autenticarse:")
    g.bullets([
        "**Cuenta guardada:** `QiskitRuntimeService.save_account(channel=\"ibm_quantum_platform\", token=getpass(...), "
        "instance=..., set_as_default=True)`. El SDK la guarda en `~/.qiskit/qiskit-ibm.json`, fuera del repositorio y "
        "**sin cifrar**: hay que protegerlo con permisos 600.",
        "**Variables de entorno**, leídas sin que la clave quede en el historial de la terminal:",
    ])
    g.code('''read -rs "QISKIT_IBM_TOKEN?IBM API key (oculta): "
echo
export QISKIT_IBM_TOKEN
read "QISKIT_IBM_INSTANCE?CRN de tu instancia Open: "
export QISKIT_IBM_INSTANCE
# ... usar el CLI ...
unset QISKIT_IBM_TOKEN QISKIT_IBM_INSTANCE''', size=8.5, caption="zsh / macOS (docs/IBM_HARDWARE.md)")
    g.p("Si falla la conexión, el CLI muestra un mensaje propio y **oculta el texto de las excepciones del SDK**, que "
        "podría contener datos de la cuenta.")

    g.h2("6.4 Qué hace el script, paso a paso")
    g.steps([
        "**Valida los argumentos:** n de 2 a 6, estado marcado binario de longitud n, shots de 1 a 8192, semilla de 0 a "
        "4294967295, iteraciones de 0 a 20 (por defecto, el óptimo del laboratorio).",
        "**Abre el servicio** (`build_service`) salvo en modo `fake`.",
        "**Construye el mismo circuito** que la web: importa `build_grover_circuit`, `optimal_iterations` y "
        "`theoretical_success` de `app.quantum.experiments`.",
        "**Elige el backend:** `FakeFez()` en modo fake; el indicado con `--backend` (comprobando que es real, operativo y "
        "con qubits suficientes); o `service.least_busy(operational=True, simulator=False, min_num_qubits=n)`.",
        "**Transpila para ese chip** con `generate_preset_pass_manager(optimization_level=3, backend=backend, "
        "seed_transpiler=seed)` y anota profundidad, puertas de 2 qubits y colocación física de los qubits.",
        "**dry-run:** se detiene aquí. No crea un Sampler ni envía nada.",
        "**fake:** ejecuta en `AerSimulator.from_backend(FakeFez())`.",
        "**ibm:** guarda los parámetros dentro del propio circuito (`metadata[\"qubitguard\"]`), envía un job con "
        "`SamplerV2(mode=backend).run([compiled], shots=shots)` e **imprime el `job_id` al instante**, antes de esperar.",
        "**Espera la cola** y lee `job.result()[0].data.meas.get_counts()` (el registro `meas` lo creó `measure_all()`).",
        "**Guarda un JSON** en `results/ibm_<fecha UTC>.json` (o `--out`) sin sobrescribir nunca un archivo existente, "
        "con versiones, backend, `job_id`, métricas, conteos, éxito medido, ideal y azar.",
    ])
    g.code('''backend = service.least_busy(operational=True, simulator=False, min_num_qubits=args.n_qubits)
compiled = generate_preset_pass_manager(
    optimization_level=3, backend=backend, seed_transpiler=args.seed
).run(circuit)
...
compiled.metadata = {**(compiled.metadata or {}), "qubitguard": dict(report)}
job = SamplerV2(mode=backend).run([compiled], shots=args.shots)
report["job_id"] = job.job_id()
print(f"job_id: {report['job_id']}", flush=True)
raw_counts = job.result()[0].data.meas.get_counts()''', size=8, caption="scripts/run_ibm_hardware.py, función run (extracto)")

    g.h2("6.5 Los modos y sus comandos")
    g.table(["Modo", "¿Red y cuenta?", "Qué hace"], [
        ["`fake` (por defecto)", "No", "Transpila para FakeFez (nivel 3) y ejecuta con su modelo de ruido."],
        ["`dry-run`", "Sí", "Elige backend y transpila, **sin enviar job**. Muestra profundidad, puertas de 2 qubits y colocación."],
        ["`ibm`", "Sí, **consume cuota**", "Misma compilación; envía un circuito con SamplerV2 e imprime el `job_id`."],
        ["`ibm --job-id ID`", "Sí", "Recupera el resultado de un job ya enviado por este CLI, sin reenviarlo, con los parámetros originales guardados en el job."],
    ], [3.4, 3, 9.6], size=9.5)
    g.code('''# Sin cuenta ni red: el mismo circuito sobre el modelo de ruido FakeFez
python -m scripts.run_ibm_hardware --mode fake --n-qubits 3 --marked 101 --shots 1024 --seed 7

# Con cuenta: revisar primero la compilación (no envía nada)
python -m scripts.run_ibm_hardware --mode dry-run --n-qubits 3 --marked 101 --seed 7

# Ejecución real (consume cuota): un solo job
python -m scripts.run_ibm_hardware --mode ibm --n-qubits 3 --marked 101 --shots 1024 --seed 7

# Recuperar un job si se cerró la terminal
python -m scripts.run_ibm_hardware --mode ibm --job-id ID_DEL_JOB --out results/recuperado.json''',
           size=8, caption="Comandos reales (desde la raíz del repositorio)")
    g.p("La semilla controla la transpilación y, en modo fake, el muestreo de Aer; **no fija las mediciones del "
        "hardware real**, que siempre tienen variación.")

    g.h2("6.6 Qué esperar: cifras medidas con FakeFez (simulación)")
    rows = [[str(n), f"`{m}`", str(k), str(d), str(q), f"{hits}/1024", f"{fmt(ms*100, 1)} %", f"{fmt(ide*100, 1)} %",
             f"{fmt(r*100, 2)} %"] for n, m, k, d, q, hits, ms, ide, r in FAKEFEZ]
    g.table(["n", "Marcado", "k", "Profundidad", "Puertas 2q", "Aciertos", "Éxito FakeFez", "Ideal", "Azar 1/N"], rows,
            [0.8, 2, 0.8, 2.2, 2, 2.2, 2.4, 1.8, 2], size=9,
            caption="CLI en modo fake, 24-10-2026 UTC, 1024 shots, seed 7, nivel 3, Qiskit 2.6.1, Aer 0.18.0, Runtime 0.51.0. "
                    "SIMULACIÓN, no resultados de una QPU (docs/IBM_HARDWARE.md). La fila n=3 se volvió a reproducir para esta guía.")
    g.figure(g.fig["noise"], 14, "Éxito ideal frente a FakeFez y frente al azar. Con n = 6 el éxito simulado (1,7 %) es "
             "indistinguible del azar (1,6 %). Datos de la tabla anterior.")
    g.p("**Lectura:** con 3 qubits la señal sigue muy por encima del azar (80 % frente a 12,5 %). Con 6 qubits, la "
        "amplificación desaparece. En un procesador real cabe esperar algo parecido: con n = 3, un éxito menor que el "
        "ideal y frecuentemente entre el ideal y el azar, **sin garantía** (ruido sesgado y fluctuaciones pueden dar "
        "incluso menos que 1/N). Backend, calibración del día, versiones, colocación y estado marcado cambian el resultado.")

    g.h2("6.7 Por qué el ruido y la profundidad degradan el resultado")
    g.bullets([
        "**Decoherencia:** un qubit real pierde su estado con el tiempo, por relajación de energía (tiempo T1) y por "
        "pérdida de la fase (tiempo T2). Cuanto más profundo es el circuito, más tiempo deben aguantar los qubits.",
        "**Error de las puertas de 2 qubits:** son mucho menos precisas que las de 1 qubit. Cada una añade un poco de "
        "error, y se acumula.",
        "**Mapeo de qubits y SWAPs:** una `mcx` de muchos controles se descompone en muchas puertas de 2 qubits, y como "
        "en el chip solo interactúan vecinos, el transpilador añade SWAPs (3 puertas de 2 qubits cada uno). Por eso las "
        "puertas de 2 qubits pasan de 39 (n = 3) a 1706 (n = 6) y la profundidad de 137 a 4447.",
        "**Error de lectura:** a veces se lee un 1 cuando era 0, o al revés.",
    ])
    g.callout("ejemplo", [
        "**Cálculo ilustrativo** (no es la calibración real del chip): si cada puerta de 2 qubits funcionara bien con "
        "probabilidad 0,995, la probabilidad de que 39 puertas salgan todas bien sería `0,995^39 ≈ 0,82`; con 1706 "
        "puertas, `0,995^1706 ≈ 0,0002`. Esa es la intuición de por qué el circuito grande «se olvida» de la respuesta "
        "y su distribución se parece a la del azar.",
    ])
    g.h2("6.8 Plan para correrlo de verdad")
    g.steps([
        "Configurar la cuenta con una de las dos formas seguras de la sección 6.3.",
        "Empezar con **n = 3** y ejecutar primero `dry-run` para ver la profundidad en el chip elegido.",
        "Enviar **un solo job** en modo `ibm` y **guardar el `job_id`**.",
        "Si se corta la conexión, recuperar con `--job-id`, sin reenviar.",
        "Citar el resultado así: «QubitGuard, backend `<backend>`, reporte `<fecha UTC>`, job `<job_id>`, n=3, "
        "marcado 101, k=2, 1024 shots; Qiskit `<versión>`, Runtime `<versión>`», adjuntando el JSON.",
    ])
    g.callout("jurado", [
        "**«¿Cómo lo correrían en hardware real?»**",
        "Con el CLI `scripts/run_ibm_hardware.py`, que usa el mismo circuito de la web. Se autentica con una cuenta Open de "
        "IBM Quantum Platform leída de variables de entorno, elige el procesador menos ocupado con `least_busy`, "
        "transpila con nivel 3 para ese chip, envía el job con `SamplerV2`, imprime el `job_id` y guarda los conteos en un "
        "JSON. Está probado en local con FakeFez y con dobles del servicio; la ejecución real está pendiente porque no "
        "teníamos cuenta configurada.",
    ])


def chapter7(g: Guide, f: dict) -> None:
    g.h1("7. Validación, pruebas y limitaciones")
    g.h2("7.1 Pruebas automáticas")
    g.p("**72 pruebas de Python (pytest)** y **37 del frontend (node:test)** pasan, además de la comprobación de tipos "
        "de TypeScript (`npm run check`). La integración continua (GitHub Actions) ejecuta Python 3.13 sin las 2 pruebas "
        "lentas de FakeFez, y Node 24 para el frontend.")
    g.table(["Archivo", "Pruebas", "Qué comprueba"], [
        ["`tests/test_quantum.py`", "20", "Iteraciones óptimas de 2 a 6 qubits, curva uniforme en k = 0 y máxima en el óptimo, conteos ideales con semilla, Grover de 2 qubits, Monte Carlo, comparación, ejecución con FakeFez (lenta) y curva Statevector igual a la fórmula en cada k."],
        ["`tests/test_api.py`", "19", "Formas de respuesta, rechazo de entradas inválidas y campos extra (422), semilla fuera de rango (422, no 500), contenido educativo y reto diario."],
        ["`tests/test_ibm_script.py`", "19", "Argumentos inválidos, rechazo de `--token`, lectura de variables de entorno, errores de cuenta sin fugas, modo fake de punta a punta, dry-run sin envío, filtros de `least_busy`, backend con nombre, envío con Sampler y recuperación con metadatos originales, no sobrescribir `--out`. Todo con dobles locales, sin red."],
        ["`tests/test_security_headers.py`", "10", "Cabeceras de seguridad, CSP de docs, caché por tipo de respuesta, gzip, copias aisladas de la caché, 503 con `Retry-After`, solo archivos públicos, intentos de path traversal y cabeceras en errores 500."],
        ["`tests/test_privacy.py`", "4", "OpenAPI sin campos de tipo contraseña, POST rechaza un campo de contraseña, los módulos de contraseñas no usan red ni almacenamiento, y el frontend no usa `Math.random`."],
        ["`frontend/test/generator.test.mjs`", "17", "Generación, aleatoriedad, entropía e inclusión-exclusión."],
        ["`frontend/test/progress.test.mjs`", "6", "Reglas de XP, racha, niveles y topes diarios."],
        ["`dom`, `ambience`, `explorer`, `password`", "4 + 4 + 3 + 3", "Interrupción de animaciones y movimiento reducido; fondo e intro; fórmulas del explorador; estimaciones del analizador."],
    ], [4.6, 1.8, 9.6], size=8.8, caption="Pruebas por archivo (conteo con pytest --collect-only y node:test).")
    g.h2("7.2 Resultados medidos que podemos citar")
    i4 = f["ideal4"]
    g.table(["Experimento", "Resultado", "Estado"], [
        ["Grover ideal, `1010`, 4 qubits, k = 3, 2048 shots, seed 7", f"Éxito {fmt(i4['measured_success'], 4)} frente a {fmt(i4['theoretical_success'], 4)} teórico; profundidad {i4['depth']} y {i4['two_qubit_gates']} puertas de 2 qubits.", "Medido."],
        ["Curva Statevector frente a la fórmula, n = 2 a 6", "Error máximo 6,15·10⁻¹⁴ (la prueba exige < 10⁻¹²).", "Medido."],
        ["Monte Carlo, 10 bits, 5000 ensayos, seed 1", f"Media {fmt(f['mc10']['mean_attempts'], 3)} frente a 512,5.", "Medido."],
        ["CLI FakeFez, n = 3 a 6", "0,802; 0,662; 0,184; 0,0166.", "Medido (simulado)."],
        ["Grover con ruido, 6 qubits, 256 shots, seed 7 (API)", "Éxito 0,03125; azar 0,0156.", "Medido (simulado)."],
        ["Ejecución en un procesador de IBM", "Ninguna.", "**No realizada.**"],
        ["Piloto con estudiantes", "Ninguno.", "**Pendiente.**"],
    ], [5.8, 7, 3.2], size=9, caption="Resultados reproducibles (README, sección 7).")
    g.h2("7.3 Limitaciones honestas")
    g.steps([
        "**El oráculo conoce la respuesta.** No se programan hashes ni cifrados para buscar secretos desconocidos.",
        "**La ventaja es cuadrática, no exponencial.** Duplicar los bits restaura el margen. Shor no se simula.",
        "**No se implementa QAE.** El Monte Carlo es búsqueda clásica de referencia; no hay machine learning.",
        "**La simulación llega a 6 qubits (64 estados).** 128 o 256 bits son extrapolaciones con fórmulas; el explorador "
        "llega a 10 qubits porque usa fórmulas, no circuitos.",
        "**El ruido es simulado** (FakeFez). Además, los modos ideal y ruidoso usan objetivos y niveles de optimización "
        "distintos, así que la diferencia no aísla solo el ruido.",
        "**Sin hardware real.** El CLI existe; la cuenta, la cola y el resultado en QPU están pendientes.",
        "**No existe hoy un ataque cuántico práctico a contraseñas ni a AES.**",
        "**El analizador de contraseñas es aproximado** y los tiempos dependen de tasas supuestas.",
        "**El progreso es local y editable:** sin cuentas, respaldo ni ranking confiable.",
        "**Capacidad del servidor:** verificado por LAN y HTTPS en CasaOS, sin pruebas de carga masiva. El experimento de "
        "6 qubits con ruido y 256 shots tardó 28,52 s en ese equipo.",
        "**Piloto con usuarios y retención:** pendientes. Protocolo propuesto: cinco participantes, tres preguntas antes y "
        "después, tiempos y dificultades registrados.",
    ])
    g.callout("idea", [
        "Las limitaciones no son una debilidad que esconder: son parte del contenido educativo. La app las muestra en su "
        "propia sección y distingue siempre entre *teórico*, *medido* e *hipotético*.",
    ])


QA = [
    ("¿Por qué usar computación cuántica para esto?",
     "Porque el concepto que enseñamos, que las amplitudes pueden ser negativas e interferir, solo existe en cuántica. "
     "Una animación con fórmulas no basta: construimos, transpilamos y medimos circuitos reales con Qiskit, y mostramos "
     "lo que pasa con el ruido de un chip de IBM simulado. Lo que no necesita cuántica (Monte Carlo, contraseñas) no la usa."),
    ("¿Esto rompe contraseñas hoy?",
     "No. El oráculo conoce la respuesta porque se la damos, el laboratorio llega a 6 qubits, y no existe hoy un ataque "
     "cuántico práctico a contraseñas ni a AES. Grover da una ventaja cuadrática que se neutraliza duplicando la longitud de "
     "la clave. El riesgo cuántico concreto es para RSA y curvas elípticas por Shor, y para eso NIST publicó FIPS 203, 204 y 205."),
    ("¿Qué es FakeFez?",
     "Un backend falso de qiskit-ibm-runtime: una copia local de las características de un procesador IBM de 156 qubits "
     "(puertas nativas, conexiones y una instantánea de sus errores). Con AerSimulator.from_backend lo usamos para "
     "simular el ruido sin cuenta. Es una simulación, no el chip."),
    ("¿Por qué no implementaron QAE, si es el par oficial de Monte Carlo?",
     "Porque nuestro eje es el otro par oficial, distribuciones de probabilidad → amplitudes, y lo cubrimos completo. "
     "Nuestro Monte Carlo es clásico y sirve de referencia para medir la ventaja de Grover. Implementar QAE sería otro "
     "proyecto, con su propio diseño y validación, y preferimos no fingirlo."),
    ("¿Desde qué edad se puede usar?",
     "Lo proponemos desde 14 años: solo exige fracciones, porcentajes y la idea de probabilidad, no álgebra lineal. Es un "
     "público propuesto; todavía no lo validamos con estudiantes de esa edad, y lo decimos."),
    ("¿Por qué es adictivo o por qué volverían?",
     "Porque está diseñado para unos 3 minutos al día: un reto diario nuevo e igual para todos, racha de días seguidos, XP, "
     "siete niveles, una ruta de cinco pasos y seis logros. Y porque cada predicción es una pequeña apuesta que se gana o "
     "se pierde. Es un diseño: su efecto en el hábito no está medido."),
    ("¿Cómo se usaría en hardware real?",
     "Con scripts/run_ibm_hardware.py: cuenta Open de IBM Quantum Platform por variables de entorno, least_busy para elegir "
     "el chip, transpilación de nivel 3, SamplerV2 para enviar el job, job_id impreso al instante y resultados en JSON. "
     "Probado con FakeFez y dobles locales; la ejecución real está pendiente."),
    ("¿La contraseña se envía al servidor?",
     "No. El analizador y el generador corren en el navegador; la API no tiene campos de contraseña, rechaza campos extra y "
     "hay pruebas que lo verifican."),
    ("¿Qué pasa si alguien hace ingeniería inversa del generador?",
     "Nada útil: el código ya es público. La seguridad no depende de ocultar el algoritmo sino de que los bits aleatorios "
     "salen de crypto.getRandomValues, el generador seguro del sistema, sin semilla, sin guardarse y sin enviarse. Conocer "
     "el código dice cómo se combinan los bits, no cuáles salieron."),
    ("¿Cómo lo validaron?",
     "72 pruebas de Python y 37 de frontend, todas pasan: física (la curva de Qiskit coincide con la fórmula con error "
     "máximo de 6,15·10⁻¹⁴), contrato de la API, privacidad, cabeceras de seguridad y el CLI de hardware con dobles. "
     "Lo que no validamos lo declaramos: hardware real y piloto con estudiantes."),
    ("¿Por qué el éxito con ruido baja tanto con 6 qubits?",
     "Porque el circuito transpilado para el chip pasa de 39 a 1706 puertas de 2 qubits y de profundidad 137 a 4447: "
     "las puertas multicontroladas se descomponen en muchas puertas pequeñas, hacen falta SWAPs por la conectividad, y "
     "cada puerta añade error. Con n = 6 el éxito simulado (1,7 %) queda al nivel del azar (1,6 %)."),
    ("¿La ventaja de Grover es exponencial?",
     "No, es cuadrática: pasa de unos N/2 intentos a unas π/4·√N iteraciones. Con 4 qubits, 8,5 intentos frente a 3 "
     "iteraciones; con 128 bits, 1,7·10^38 frente a 1,45·10^19. Grande, pero no vuelve fácil un problema enorme."),
    ("¿Por qué el explorador llega a 10 qubits y el laboratorio solo a 6?",
     "El laboratorio ejecuta circuitos en el servidor y debe responder rápido; el explorador usa las fórmulas exactas en el "
     "navegador y responde al instante. Sus valores coinciden con Statevector de 2 a 6 qubits."),
    ("¿Qué es transpilar y por qué importa?",
     "Reescribir el circuito con las puertas nativas y las conexiones del chip, y optimizarlo. Importa porque la "
     "profundidad y el número de puertas de 2 qubits tras transpilar determinan cuánto ruido acumula el circuito en "
     "hardware real."),
    ("¿Los resultados son reproducibles?",
     "Sí, dentro del mismo entorno: la API acepta una semilla que fija transpilación y muestreo, y las cifras citadas "
     "indican versión, semilla y shots. En hardware real la semilla no fija las mediciones."),
    ("¿Qué harían con más tiempo?",
     "El piloto con cinco estudiantes con preguntas antes y después, una ejecución real en IBM con n = 3 empezando por "
     "dry-run, y medir la retención del reto diario. Una lección de QAE sería un alcance nuevo."),
]


def chapter8(g: Guide) -> None:
    g.h1("8. Preguntas probables del jurado")
    g.h2("8.1 El resumen de 30 segundos")
    g.callout("idea", [
        "«QubitGuard es un laboratorio web para aprender probabilidad cuántica. El estudiante predice cuánto acertará el "
        "algoritmo de Grover, ejecuta un circuito real de Qiskit en un simulador, y compara su predicción con la teoría y "
        "con lo medido, también con el ruido simulado de un chip de IBM. Enseña que en cuántica se manipulan amplitudes, "
        "que pueden ser negativas, y desmonta el mito de que las computadoras cuánticas ya rompen contraseñas. Todo está "
        "probado; no lo ejecutamos todavía en hardware real, pero el script para hacerlo está listo.»",
    ])
    g.h2("8.2 Respuestas modelo")
    g.p("Cada respuesta cabe en 20 a 40 segundos. Apréndete la idea, no la letra.")
    for i, (q, a) in enumerate(QA, 1):
        para = g.p(f"{i}. {q}", keep=True, after=2)
        for run in para.runs:
            run.bold = True
            run.font.color.rgb = H2_COLOR
        ans = g.p(a, after=9)
        ans.paragraph_format.left_indent = Cm(0.6)
    g.h2("8.3 Frases que nunca debemos decir")
    g.table(["No digas", "Di en su lugar"], [
        ["«Lo corrimos en una computadora cuántica de IBM.»", "«Lo corrimos en simuladores, también con el modelo de ruido FakeFez; el script para hardware real está listo.»"],
        ["«Implementamos QAE.»", "«Cubrimos el par distribuciones → amplitudes; nuestro Monte Carlo es clásico.»"],
        ["«Grover rompe contraseñas.»", "«Grover acelera de forma cuadrática una búsqueda; duplicar la clave restaura el margen.»"],
        ["«Prueba todas las respuestas a la vez.»", "«Usa interferencia para que la respuesta correcta sea la más probable.»"],
        ["«Está probado que los alumnos aprenden más.»", "«El piloto con estudiantes está pendiente.»"],
        ["«Tu contraseña es segura contra Grover.»", "«Supera la meta educativa de 128 bits; no es una certificación.»"],
    ], [7, 9], size=9.5)


GLOSSARY = [
    ("Amplitud", "Número asociado a cada resultado posible de un estado cuántico; su cuadrado es la probabilidad. Puede ser negativo."),
    ("AerSimulator", "Simulador de Qiskit Aer que ejecuta circuitos con shots y devuelve conteos; admite modelos de ruido."),
    ("Backend", "Destino donde se ejecuta un circuito: un simulador o un procesador cuántico."),
    ("Bit", "Unidad de información clásica: vale 0 o 1."),
    ("Circuito cuántico", "Secuencia de puertas aplicadas a qubits, leída de izquierda a derecha, que termina en mediciones."),
    ("Counts (conteos)", "Cuántas veces salió cada cadena de bits al repetir un circuito muchas veces."),
    ("CSPRNG", "Generador de números aleatorios criptográficamente seguro; en el navegador, crypto.getRandomValues."),
    ("CZ", "Puerta de 2 qubits que cambia el signo de |11⟩; puerta nativa de 2 qubits de FakeFez."),
    ("Decoherencia", "Pérdida del estado cuántico de un qubit con el tiempo por su interacción con el entorno (tiempos T1 y T2)."),
    ("Difusor", "Bloque de Grover que refleja todas las amplitudes respecto a su media (inversión sobre la media)."),
    ("Entropía", "Medida, en bits, de cuántas posibilidades igual de probables tiene un secreto: log2 del número de posibilidades."),
    ("Estimator", "Primitiva de Qiskit Runtime que calcula valores esperados de observables. No se usa en QubitGuard."),
    ("Fake backend", "Objeto local que imita un procesador real de IBM (qubits, puertas, conexiones y ruido) sin conectarse a IBM."),
    ("FakeFez", "Fake backend de qiskit-ibm-runtime que imita a ibm_fez, un procesador de 156 qubits."),
    ("FIPS 203, 204, 205", "Estándares de criptografía poscuántica de NIST (2024): ML-KEM, ML-DSA y SLH-DSA."),
    ("Grover", "Algoritmo cuántico de búsqueda sin estructura que necesita unas π/4·√N iteraciones en lugar de unos N/2 intentos."),
    ("H (Hadamard)", "Puerta de 1 qubit que crea o deshace superposición."),
    ("Interferencia", "Suma de amplitudes: con igual signo se refuerzan, con signo opuesto se cancelan."),
    ("Iteración de Grover", "Una aplicación del oráculo seguida del difusor; gira el estado 2θ hacia el marcado."),
    ("Job", "Trabajo enviado a IBM Quantum; espera en cola, se ejecuta y queda identificado por su job_id."),
    ("Layout", "Asignación de qubits lógicos del circuito a qubits físicos del chip."),
    ("least_busy", "Método de QiskitRuntimeService que elige el procesador con menos cola que cumpla unos filtros."),
    ("Mapa de conexiones", "Qué pares de qubits físicos pueden interactuar directamente en un chip."),
    ("MCX / MCZ", "Puertas X y Z multicontroladas: actúan solo si todos los controles valen 1."),
    ("Medición", "Operación que convierte el estado cuántico en un resultado clásico y destruye la superposición."),
    ("Monte Carlo", "Método que estima resultados repitiendo experimentos aleatorios. En QubitGuard, la búsqueda clásica al azar."),
    ("NISQ", "Era actual de procesadores cuánticos ruidosos de escala intermedia, sin corrección de errores completa."),
    ("Oráculo de fase", "Bloque que cambia el signo de la amplitud del estado marcado."),
    ("Pass manager", "Tubería de pasadas de transpilación (layout, routing, traducción, optimización)."),
    ("POE", "Predice, Observa, Explica: metodología en la que el estudiante se compromete con una predicción antes de ver el resultado."),
    ("Primitiva", "Interfaz estándar de Qiskit Runtime para ejecutar circuitos: Sampler o Estimator."),
    ("Profundidad", "Número de capas de puertas que no pueden ejecutarse en paralelo; mide el «largo» del circuito."),
    ("QAE", "Quantum Amplitude Estimation: algoritmo que estima una probabilidad con menos muestras que Monte Carlo. No implementado."),
    ("QPU", "Unidad de procesamiento cuántico: el chip cuántico real."),
    ("Qiskit", "SDK de código abierto de IBM para construir, transpilar y ejecutar circuitos cuánticos."),
    ("Qiskit Runtime", "Servicio y paquete (qiskit-ibm-runtime) para ejecutar en procesadores de IBM mediante primitivas."),
    ("Qubit", "Unidad de información cuántica; su estado se describe con amplitudes para 0 y 1."),
    ("Sampler / SamplerV2", "Primitiva que ejecuta circuitos y devuelve muestras de las cadenas de bits medidas."),
    ("Semilla (seed)", "Número que fija un generador pseudoaleatorio para que los resultados se puedan repetir."),
    ("Shor", "Algoritmo cuántico que factoriza enteros eficientemente; amenaza a RSA y curvas elípticas. No se simula aquí."),
    ("Shot", "Una repetición de preparar y medir un circuito."),
    ("Statevector", "Clase de Qiskit que calcula el estado exacto de un circuito sin medir."),
    ("Superposición", "Estado con amplitudes no nulas para varios resultados; al medir se obtiene uno solo."),
    ("SWAP", "Operación que intercambia el estado de dos qubits; cuesta 3 puertas de 2 qubits."),
    ("Transpilar", "Reescribir y optimizar un circuito para las puertas nativas y las conexiones de un backend."),
]

REFERENCES = [
    "L. K. Grover, «A fast quantum mechanical algorithm for database search», STOC 1996. arXiv:quant-ph/9605043.",
    "M. Boyer, G. Brassard, P. Høyer y A. Tapp, «Tight bounds on quantum searching», Fortschritte der Physik 46, 1998. arXiv:quant-ph/9605034.",
    "P. W. Shor, «Polynomial-time algorithms for prime factorization and discrete logarithms on a quantum computer», SIAM J. Comput. 26, 1997. arXiv:quant-ph/9508027.",
    "M. A. Nielsen e I. L. Chuang, *Quantum Computation and Quantum Information*, Cambridge University Press (edición del 10.º aniversario, 2010). Capítulo 6: algoritmos de búsqueda cuántica.",
    "Documentación de Qiskit e IBM Quantum: docs.quantum.ibm.com (redirige a quantum.cloud.ibm.com/docs). Guías de transpilación con preset pass managers, SamplerV2 y planes.",
    "Código fuente de Qiskit: github.com/Qiskit/qiskit; Qiskit Aer: github.com/Qiskit/qiskit-aer; qiskit-ibm-runtime: github.com/Qiskit/qiskit-ibm-runtime.",
    "Presentación de Qiskit: ibm.com/quantum/qiskit.",
    "NIST, FIPS 203 (ML-KEM), FIPS 204 (ML-DSA) y FIPS 205 (SLH-DSA), 2024. csrc.nist.gov/projects/post-quantum-cryptography.",
    "Páginas del evento: ibmquantumhackathon.ieeecspucp.com (retos) y qiskitfallfest2026.ieeecspucp.com/resources.",
    "Repositorio del proyecto: github.com/LeonardoShande/QubitGuard. Documentos internos citados: README.md, docs/GUIA_DEL_CODIGO.md, docs/API.md, docs/IBM_HARDWARE.md, docs/LIMITACIONES.md, docs/architecture.md, docs/RETO_Y_ENTREGA.md, pitch/OUTLINE_PRESENTACION.md.",
]


def chapter9(g: Guide) -> None:
    g.h1("9. Glosario y referencias")
    g.h2("9.1 Glosario")
    g.table(["Término", "Significado"], [[f"**{t}**", d] for t, d in GLOSSARY], [3.8, 12.2], size=9.2)
    g.h2("9.2 Referencias")
    for i, ref in enumerate(REFERENCES, 1):
        para = g.p(f"[{i}] {ref}", size=10, after=4)
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.first_line_indent = Cm(-0.8)
    g.p("Las tarjetas poscuánticas de la aplicación enlazan sus propias fuentes en `app/content/pqc.json`.", size=10,
        color=MUTED)


def build() -> Path:
    facts = qiskit_facts()
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        figs = {"geometry": tmpdir / "geometry.png", "curve": tmpdir / "curve.png", "noise": tmpdir / "noise.png"}
        fig_geometry(figs["geometry"])
        fig_curve(figs["curve"], facts["curve4"])
        fig_noise(figs["noise"])

        g = Guide()
        g.fig = figs
        cover(g, facts)
        how_to_use(g)
        chapter1(g, facts)
        chapter2(g, facts)
        chapter3(g, facts)
        chapter4(g, facts)
        chapter5(g, facts)
        chapter6(g, facts)
        chapter7(g, facts)
        chapter8(g)
        chapter9(g)
        g.insert_toc(g.toc_anchor)
        props = g.doc.core_properties
        props.title = "QubitGuard: guía de estudio del equipo"
        props.subject = "IBM Quantum Hackathon / Qiskit Fall Fest 2026, reto 5.3 EdTech"
        props.author = "Equipo QubitGuard"
        props.language = "es-PE"
        g.doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
