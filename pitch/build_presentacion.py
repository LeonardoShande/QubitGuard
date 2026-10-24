"""Genera pitch/QubitGuard_presentacion.pptx: presentación detallada, 16 slides 16:9.

Sigue pitch/OUTLINE_PRESENTACION.md (títulos, orden, expositor y cifras).
Recorta las capturas reales de pitch/capturas/ y dibuja dos gráficos con matplotlib
en pitch/capturas/derivadas/. Reproducible:

    .venv/bin/python pitch/build_presentacion.py

Requiere python-pptx 1.0.2, Pillow y matplotlib (no son dependencias de la app).
"""
from __future__ import annotations

import math
import os
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "capturas"
DERIV = SHOTS / "derivadas"
OUT = ROOT / "QubitGuard_presentacion.pptx"
TOTAL = 16

# Tokens de la app (frontend/styles.css): cian = cuántico, magenta = clásico.
HEX = {
    "bg": "080c14", "panel": "0e1520", "line": "253246", "text": "dce5ee",
    "muted": "a1b0c2", "cyan": "46d9e8", "mag": "ea6bbd", "amber": "f0b357",
    "ink": "06222a",
}
C = {k: RGBColor.from_string(v) for k, v in HEX.items()}
# QG_MONO solo para QA: LibreOffice sustituye Consolas por una serif; DejaVu Sans Mono es más ancha (holgura).
SANS, MONO = "Calibri", os.environ.get("QG_MONO", "Consolas")
BRAND = "QubitGuard · Qiskit Fall Fest 2026"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- imágenes
def crop(name: str, box: tuple[int, int, int, int], out: str) -> Path:
    """Recorta una captura (coordenadas en píxeles de la imagen original)."""
    DERIV.mkdir(parents=True, exist_ok=True)
    dst = DERIV / out
    Image.open(SHOTS / name).crop(box).save(dst, optimize=True)
    return dst


IMG = {
    "intro": crop("01_intro.png", (560, 300, 2330, 1480), "s01_intro_ascii.png"),
    "inicio": crop("02_inicio.png", (0, 0, 2880, 1660), "s08_inicio.png"),
    "problema": crop("03_problema.png", (89, 368, 1426, 1180), "s02_problema.png"),
    "prediccion": crop("04_laboratorio_prediccion.png", (110, 1760, 750, 2545), "s09_prediccion.png"),
    "poe": crop("04_laboratorio_resultado.png", (790, 495, 2475, 1310), "s09_poe.png"),
    "curva": crop("04_laboratorio_resultado.png", (790, 2515, 2475, 3340), "s07_curva.png"),
    "circuito": crop("05_circuito.png", (40, 58, 1580, 400), "s06_circuito.png"),
    "explorador": crop("07_explorador.png", (728, 775, 2452, 1728), "s10_explorador.png"),
    "ruido": crop("08_ruido.png", (795, 1320, 2480, 1950), "s11_histograma_ruido.png"),
    "generador": crop("09_generador.png", (1215, 2390, 2420, 2815), "s13_generador.png"),
    "privacidad": crop("09_generador.png", (85, 470, 2460, 580), "s14_privacidad.png"),
    "pqc": crop("10_pqc.png", (1688, 378, 2454, 640), "s13_pqc_nist.png"),
    "reto": crop("11_reto_diario.png", (89, 425, 2452, 1228), "s16_reto_diario.png"),
}


# ---------------------------------------------------------------- gráficos
def _register_fonts() -> tuple[str, str]:
    lo = Path("/Applications/LibreOffice.app/Contents/Resources/fonts/truetype")
    sans, mono = "DejaVu Sans", "DejaVu Sans Mono"
    for f, fam in (("Carlito-Regular.ttf", "sans"), ("Carlito-Bold.ttf", None),
                   ("DejaVuSansMono.ttf", "mono"), ("DejaVuSansMono-Bold.ttf", None)):
        p = lo / f
        if p.exists():
            font_manager.fontManager.addfont(str(p))
            if fam == "sans":
                sans = "Carlito"
    return sans, mono


def _style(sans: str) -> None:
    plt.rcParams.update({
        "font.family": sans, "font.size": 15, "text.color": "#" + HEX["text"],
        "axes.facecolor": "#" + HEX["bg"], "figure.facecolor": "#" + HEX["bg"],
        "axes.edgecolor": "#" + HEX["line"], "axes.labelcolor": "#" + HEX["muted"],
        "xtick.color": "#" + HEX["muted"], "ytick.color": "#" + HEX["muted"],
        "axes.grid": True, "grid.color": "#18222f", "grid.linewidth": 1,
        "mathtext.fontset": "custom", "mathtext.rm": sans,
    })


def _es(x: float, nd: int) -> str:
    return f"{x:.{nd}f}".replace(".", ",")


def chart_busqueda(sans: str, mono: str) -> Path:
    """Clásico (N+1)/2 frente a Grover π/4·√N, escala logarítmica."""
    cyan, mag, muted = "#" + HEX["cyan"], "#" + HEX["mag"], "#" + HEX["muted"]
    n = list(range(1, 41))
    clas = [(2 ** b + 1) / 2 for b in n]
    grov = [math.pi / 4 * math.sqrt(2 ** b) for b in n]
    fig, ax = plt.subplots(figsize=(7.4, 4.75), dpi=220)
    ax.semilogy(n, clas, color=mag, lw=2.8)
    ax.semilogy(n, grov, color=cyan, lw=2.8)
    ax.axvline(4, color=muted, lw=1, ls=":")
    ax.plot([4], [8.5], "o", color=mag, ms=8)
    ax.plot([4], [math.pi], "o", color=cyan, ms=8)
    ax.annotate("8,5 intentos", (4, 8.5), (5.2, 2e4), color=mag, fontsize=15,
                fontfamily=mono, arrowprops=dict(arrowstyle="-", color=mag, lw=1))
    ax.annotate("3 iteraciones", (4, math.pi), (7, 0.5), color=cyan, fontsize=15,
                fontfamily=mono, arrowprops=dict(arrowstyle="-", color=cyan, lw=1))
    ax.text(4.4, 3e10, "N = 16", color=muted, fontsize=14, fontfamily=mono)
    ax.text(13, 3e9, "Al azar  (N+1)/2", color=mag, fontsize=16, weight="bold", rotation=0)
    ax.text(25, 6e1, "Grover  π/4·√N", color=cyan, fontsize=16, weight="bold")
    ax.set_yticks([1, 1e3, 1e6, 1e9, 1e12])
    ax.set_yticklabels([r"$10^{0}$", r"$10^{3}$", r"$10^{6}$", r"$10^{9}$", r"$10^{12}$"])
    ax.set_ylim(0.3, 3e12)
    ax.set_xlim(1, 40)
    ax.set_xticks([1, 4, 10, 20, 30, 40])
    ax.set_xlabel(r"bits del secreto  n     (N = $2^n$ posibilidades)")
    ax.set_ylabel("consultas necesarias")
    ax.grid(axis="x", visible=False)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    dst = DERIV / "s05_grafico_azar_vs_grover.png"
    fig.savefig(dst, facecolor=fig.get_facecolor())
    plt.close(fig)
    return dst


def chart_ruido(sans: str, mono: str) -> Path:
    """Éxito FakeFez (CLI, seed 7, 1024 shots) frente al azar 1/N, n=3..6."""
    cyan, mag = "#" + HEX["cyan"], "#" + HEX["mag"]
    ns = [3, 4, 5, 6]
    medido = [0.802, 0.662, 0.184, 0.0166]
    azar = [0.125, 0.0625, 0.031, 0.0156]
    prof = [137, 448, 1456, 4447]
    fig, ax = plt.subplots(figsize=(7.6, 4.75), dpi=220)
    w, off = 0.34, 0.2
    xs = range(len(ns))
    ax.bar([x - off for x in xs], medido, w, color=cyan, label="FakeFez medido")
    ax.bar([x + off for x in xs], azar, w, color=mag, label="Azar 1/N")
    for x, m, a in zip(xs, medido, azar):
        lift = 0.07 if m < 0.05 else 0  # n=6: las dos barras casi iguales, separa etiquetas
        ax.text(x - off, m + 0.02 + lift, _es(m * 100, 1) + " %", ha="center", color=cyan,
                fontsize=12.5, fontfamily=mono, weight="bold")
        ax.text(x + off, a + 0.02, _es(a * 100, 1) + " %", ha="center", color=mag,
                fontsize=12.5, fontfamily=mono)
    ax.set_xticks(list(xs))
    ax.set_xticklabels([f"n = {n}\nprofundidad {p}" for n, p in zip(ns, prof)], fontsize=14)
    ax.set_ylim(0, 1)
    ax.set_yticks([0, 0.25, 0.5, 0.75, 1])
    ax.set_yticklabels(["0 %", "25 %", "50 %", "75 %", "100 %"])
    ax.set_ylabel("probabilidad de acertar")
    ax.grid(axis="x", visible=False)
    ax.legend(frameon=False, fontsize=15, loc="upper right", labelcolor="#" + HEX["text"])
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    dst = DERIV / "s11_grafico_ruido_fakefez.png"
    fig.savefig(dst, facecolor=fig.get_facecolor())
    plt.close(fig)
    return dst


_SANS, _MONO = _register_fonts()
_style(_SANS)
IMG["g_busqueda"] = chart_busqueda(_SANS, _MONO)
IMG["g_ruido"] = chart_ruido(_SANS, _MONO)


# ---------------------------------------------------------------- primitivas
MARK = re.compile(r"(\[\[.+?\]\]|<<.+?>>|\{\{.+?\}\})")


def _runs(p, line: str, size: float, color: RGBColor, bold: bool, font: str) -> None:
    """[[x]] cian negrita (cuántico), <<x>> magenta negrita (clásico), {{x}} monoespaciada."""
    for tok in MARK.split(line):
        if not tok:
            continue
        r = p.add_run()
        f, col, b = font, color, bold
        if tok.startswith("[["):
            tok, col, b = tok[2:-2], C["cyan"], True
        elif tok.startswith("<<"):
            tok, col, b = tok[2:-2], C["mag"], True
        elif tok.startswith("{{"):
            tok, f = tok[2:-2], MONO
        r.text = tok
        r.font.size, r.font.bold, r.font.name = Pt(size), b, f
        r.font.color.rgb = col


def text(s, x, y, w, h, content: str, size=18, color=None, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.08, after=0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        _runs(p, line, size, color or C["text"], bold, font)
    return tb


def panel(s, x, y, w, h, fill="panel", line="line", radius=0.12, dash=False, lw=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C[fill]
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C[line]
        shp.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rect(s, x, y, w, h, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = C[fill]
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def arrow(s, x, y, w=0.42, h=0.34, color="muted", kind=MSO_SHAPE.RIGHT_ARROW):
    shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = C[color]
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def chip_w(word: str, size=15) -> float:
    return len(word) * size * 0.62 / 72 + 0.34


def chips(s, x, y, words, accent="cyan", size=15, gap=0.14) -> float:
    for wd in words:
        w = chip_w(wd, size)
        panel(s, x, y, w, 0.42, fill="panel", line=accent, radius=0.21)
        text(s, x, y, w, 0.42, wd, size, C[accent], font=MONO, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        x += w + gap
    return x


def picture(s, path: Path, x, y, w=None, h=None, border=True):
    iw, ih = Image.open(path).size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = C["line"]
        pic.line.width = Pt(1)
    return pic, w, h


def keyrow(s, x, y, w, content, size=18, marker="cyan", h=0.9):
    rect(s, x, y + size / 72 * 0.42, 0.11, 0.11, marker)
    text(s, x + 0.3, y, w - 0.3, h, content, size)


SECTIONS = {
    1: "Contexto", 2: "Contexto", 3: "Contexto",
    4: "Qiskit y física", 5: "Qiskit y física", 6: "Qiskit y física", 7: "Qiskit y física",
    8: "La aplicación", 9: "La aplicación", 10: "La aplicación", 11: "La aplicación",
    12: "Hardware, seguridad y validación", 13: "Hardware, seguridad y validación",
    14: "Hardware, seguridad y validación", 15: "Hardware, seguridad y validación",
    16: "Cierre",
}


def slide(n: int, title: str | None, notes: str):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["bg"]
    if title:
        text(s, 0.6, 0.42, 9, 0.35, f"|{n:02d}>  {SECTIONS[n]}", 16, C["cyan"], font=MONO)
        text(s, 0.6, 0.78, 12.1, 0.8, title, 36, C["text"], bold=True)
    text(s, 0.6, 7.0, 7, 0.3, BRAND, 14, C["muted"])
    text(s, 10.73, 7.0, 2.0, 0.3, f"{n:02d} / {TOTAL}", 14, C["muted"], font=MONO,
         align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


# ================================================================ 1 portada
s = slide(1, None, "Expone A. Idea clave: QubitGuard es un laboratorio para aprender "
          "probabilidad cuántica prediciendo y ejecutando circuitos Qiskit.\n"
          "Presentar equipo, reto 5.3 EdTech y la pregunta guía: ¿cuántos intentos hacen "
          "falta para encontrar un secreto?")
picture(s, IMG["intro"], 6.55, 1.05, w=6.2, border=False)
text(s, 0.6, 1.0, 6, 0.4, "Reto 5.3 EdTech, IEEE CS PUCP", 18, C["cyan"], font=MONO)
text(s, 0.6, 1.6, 6.2, 1.2, "QubitGuard", 64, C["text"], bold=True, font=MONO)
text(s, 0.6, 2.95, 5.7, 1.6, "Un laboratorio para aprender [[probabilidad cuántica]] "
     "prediciendo y ejecutando circuitos [[Qiskit]]", 26, spacing=1.05)
chips(s, 0.6, 4.65, ["EdTech", "Qiskit", "Grover"], size=16)
text(s, 0.6, 5.55, 6, 0.4, "IBM Quantum Hackathon / Qiskit Fall Fest 2026, Lima", 18, C["muted"])
text(s, 0.6, 6.05, 6, 0.45, "Equipo: ____", 22, bold=True)

# ================================================================ 2 reto y problema
s = slide(2, "El reto y el problema",
          "Expone A. Idea clave: la probabilidad cuántica no se intuye y la amenaza cuántica "
          "se explica con alarmismo.\nPúblico: desde 14 años, unos 3 minutos al día. "
          "La captura muestra cómo la app plantea problema y solución.")
for i, (lab, col, head, body) in enumerate([
    ("Problema 1: física", "cyan", "La probabilidad cuántica no se intuye",
     "Las [[amplitudes]] pueden ser negativas y cancelarse. La estadística escolar no lo anticipa."),
    ("Problema 2: seguridad", "mag", "La amenaza se explica con alarmismo",
     "«Romperá todas las contraseñas», con fórmulas que nadie comprueba."),
]):
    y = 1.9 + i * 2.05
    panel(s, 0.6, y, 5.75, 1.85)
    text(s, 0.9, y + 0.2, 5.2, 0.3, lab, 15, C[col], font=MONO)
    text(s, 0.9, y + 0.55, 5.2, 0.45, head, 22, bold=True)
    text(s, 0.9, y + 1.02, 5.2, 0.8, body, 18, C["muted"])
chips(s, 0.6, 6.15, ["Reto 5.3", "intuición", "criptografía"])
pic, pw, ph = picture(s, IMG["problema"], 6.75, 1.9, w=5.98)
text(s, 6.75, 1.9 + ph + 0.1, 5.98, 0.35, "Sección «El problema» de la app", 14, C["muted"])

# ================================================================ 3 mínimo cuántico
s = slide(3, "Lo mínimo de computación cuántica",
          "Expone A. Idea clave: un qubit está en superposición; al medir sale 0 o 1 con "
          "probabilidad = amplitud al cuadrado.\nEjemplo: 0,6 y 0,8 dan 36 % y 64 %. "
          "Las amplitudes pueden cancelarse: eso es lo que Grover aprovecha.")
cols = [(0.6, 3.85), (4.92, 3.85), (9.24, 3.49)]
base, sc = 4.7, 1.7
# P1 amplitudes
x, w = cols[0]
panel(s, x, 1.9, w, 3.75)
text(s, x + 0.3, 2.05, w - 0.6, 0.4, "Antes de medir", 20, bold=True)
text(s, x + 0.3, 2.5, w - 0.6, 0.4, "0,6|0> + 0,8|1>", 18, C["cyan"], font=MONO)
for j, (v, lab) in enumerate([(0.6, "|0>"), (0.8, "|1>")]):
    bx = x + 0.75 + j * 1.45
    rect(s, bx, base - v * sc, 0.8, v * sc, "cyan")
    text(s, bx, base - v * sc - 0.38, 0.8, 0.32, _es(v, 1), 16, C["cyan"], font=MONO, align=PP_ALIGN.CENTER)
    text(s, bx, base + 0.1, 0.8, 0.32, lab, 16, C["muted"], font=MONO, align=PP_ALIGN.CENTER)
text(s, x + 0.3, 5.2, w - 0.6, 0.35, "amplitudes", 15, C["muted"], align=PP_ALIGN.CENTER)
arrow(s, 4.5, 3.6)
# P2 probabilidades
x, w = cols[1]
panel(s, x, 1.9, w, 3.75)
text(s, x + 0.3, 2.05, w - 0.6, 0.4, "Al medir", 20, bold=True)
text(s, x + 0.3, 2.5, w - 0.6, 0.4, "P = amplitud²", 18, C["text"], font=MONO)
for j, (v, lab) in enumerate([(0.36, "0"), (0.64, "1")]):
    bx = x + 0.75 + j * 1.45
    rect(s, bx, base - v * sc, 0.8, v * sc, "muted")
    text(s, bx - 0.1, base - v * sc - 0.38, 1.0, 0.32, f"{int(v * 100)} %", 16, C["text"], font=MONO, align=PP_ALIGN.CENTER)
    text(s, bx, base + 0.1, 0.8, 0.32, lab, 16, C["muted"], font=MONO, align=PP_ALIGN.CENTER)
text(s, x + 0.3, 5.2, w - 0.6, 0.35, "sale 0 o 1, nunca ambos", 15, C["muted"], align=PP_ALIGN.CENTER)
# P3 interferencia
x, w = cols[2]
panel(s, x, 1.9, w, 3.75)
text(s, x + 0.3, 2.05, w - 0.6, 0.4, "Interferencia", 20, bold=True)
text(s, x + 0.3, 2.5, w - 0.6, 0.4, "+0,5 y −0,5 suman 0", 18, C["cyan"], font=MONO)
mid = 4.15
s.shapes.add_connector(1, Inches(x + 0.35), Inches(mid), Inches(x + w - 0.35), Inches(mid)).line.color.rgb = C["line"]
rect(s, x + 0.75, mid - 0.5 * sc * 0.9, 0.8, 0.5 * sc * 0.9, "cyan")
rect(s, x + 1.95, mid, 0.8, 0.5 * sc * 0.9, "cyan")
text(s, x + 0.3, 5.2, w - 0.6, 0.35, "se cancelan", 15, C["muted"], align=PP_ALIGN.CENTER)
text(s, 0.6, 5.98, 6.2, 0.45, "[[n qubits]]: {{2ⁿ}} estados en [[superposición]]. 4 qubits, N = 16.", 18)
chips(s, 12.73 - sum(chip_w(k, 14) + 0.14 for k in ["qubit", "superposición", "amplitud", "medición"]) + 0.14, 5.98, ["qubit", "superposición", "amplitud", "medición"], size=14)

# ================================================================ 4 qué es Qiskit
s = slide(4, "¿Qué es Qiskit?",
          "Expone B. Idea clave: Qiskit es el kit de IBM, de código abierto, para construir, "
          "optimizar y ejecutar circuitos.\nRecorrer el flujo de 4 pasos y aterrizarlo en "
          "QubitGuard: n=4, estado 1010, k=3, 2048 shots, 0,953 medido frente a 0,961 teórico.")
text(s, 0.6, 1.75, 12, 0.45, "El [[SDK]] de IBM, de código abierto, para programar computadoras cuánticas en Python.", 20, C["muted"])
steps = [
    ("1", "Construir", "QuantumCircuit", "puertas H, oráculo y difusor"),
    ("2", "Transpilar", "transpile", "adapta el circuito a las puertas del chip"),
    ("3", "Ejecutar", "AerSimulator\nSamplerV2", "simulador local o QPU de IBM"),
    ("4", "Medir", "counts", "cuántas veces salió cada estado"),
]
bw, gap = 2.65, 0.51
for i, (num, head, code, body) in enumerate(steps):
    x = 0.6 + i * (bw + gap)
    panel(s, x, 2.4, bw, 2.2)
    text(s, x + 0.25, 2.55, 0.5, 0.4, num, 20, C["cyan"], bold=True, font=MONO)
    text(s, x + 0.65, 2.53, bw - 0.8, 0.45, head, 22, bold=True)
    text(s, x + 0.25, 3.05, bw - 0.4, 0.7, code, 16, C["cyan"], font=MONO)
    text(s, x + 0.25, 3.78, bw - 0.4, 0.8, body, 16, C["muted"])
    if i < 3:
        arrow(s, x + bw + 0.06, 3.33, w=0.39)
panel(s, 0.6, 4.85, 12.13, 1.15)
text(s, 0.9, 4.97, 3.5, 0.35, "En QubitGuard", 15, C["muted"], font=MONO)
text(s, 0.9, 5.35, 5.6, 0.6, "n=4, estado 1010, k=3, 2048 shots", 18)
text(s, 6.6, 4.97, 2.6, 0.9, "[[0,953]]", 40, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.85, 4.97, 3.7, 0.9, "medido, frente a\n0,961 teórico", 18, C["muted"], anchor=MSO_ANCHOR.MIDDLE)
chips(s, 0.6, 6.2, ["SDK", "circuito", "transpilar", "Aer", "Runtime"], size=14)

# ================================================================ 5 Monte Carlo a Grover
s = slide(5, "De Monte Carlo a Grover",
          "Expone B. Idea clave: adivinar al azar cuesta (N+1)/2 intentos; Grover necesita "
          "unas π/4·√N iteraciones.\nEjemplo N=16: 8,5 frente a 3. Monte Carlo real: 10 bits, "
          "5000 ensayos, media 513,3 frente a 512,5 esperado.")
picture(s, IMG["g_busqueda"], 0.45, 1.75, w=7.55, border=False)
x = 8.45
text(s, x, 1.85, 4.3, 0.4, "Ejemplo: N = 16", 20, bold=True)
text(s, x, 2.3, 1.7, 0.95, "8,5", 54, C["mag"], bold=True, font=MONO)
text(s, x + 1.85, 2.42, 2.45, 0.8, "intentos <<al azar>>\n{{(N+1)/2}}", 17, C["muted"])
text(s, x, 3.3, 1.7, 0.95, "3", 54, C["cyan"], bold=True, font=MONO)
text(s, x + 1.85, 3.42, 2.45, 0.8, "iteraciones de [[Grover]]\n{{π/4·√16 ≈ 3}}", 17, C["muted"])
text(s, x, 4.45, 4.28, 0.95, "[[Ventaja cuadrática]]: cuadruplicar N solo duplica el trabajo de Grover.", 18)
panel(s, x, 5.45, 4.28, 1.0)
text(s, x + 0.22, 5.55, 3.9, 0.85, "<<Monte Carlo>> medido: 10 bits, 5000 ensayos, media 513,3 (esperado 512,5)", 15, C["muted"])

# ================================================================ 6 Grover paso a paso
s = slide(6, "Grover paso a paso",
          "Expone B. Idea clave: superposición, oráculo que marca, difusor que amplifica; "
          "se repite k veces.\nEjemplo mínimo: con 2 qubits basta 1 iteración y se acierta "
          "el 100 %. La captura es el circuito real que dibuja Qiskit.")
blocks = [("H", "superposición", 0.6, 2.2), ("Oráculo", "marca: cambia el signo", 3.45, 2.75),
          ("Difusor", "amplifica: refleja respecto a la media", 6.85, 3.1), ("Medir", "counts", 10.95, 1.78)]
panel(s, 3.25, 1.78, 6.9, 1.62, fill=None, line="cyan", dash=True)
text(s, 8.65, 1.84, 1.4, 0.3, "× k veces", 15, C["cyan"], font=MONO, align=PP_ALIGN.RIGHT)
for head, body, x, w in blocks:
    panel(s, x, 2.13, w, 1.12)
    text(s, x + 0.2, 2.2, w - 0.4, 0.42, head, 20, C["cyan"] if head != "Medir" else C["text"], bold=True, font=MONO)
    text(s, x + 0.2, 2.63, w - 0.4, 0.6, body, 15, C["muted"])
arrow(s, 2.88, 2.52, w=0.3)
arrow(s, 6.28, 2.52, w=0.5)
arrow(s, 10.38, 2.52, w=0.5)
pic, pw, ph = picture(s, IMG["circuito"], 0.6, 3.75, w=8.3)
text(s, 0.6, 3.75 + ph + 0.1, 8.3, 0.35, "Circuito real dibujado por Qiskit: n=4, estado 1010, k=3 (primer tramo)", 14, C["muted"])
panel(s, 9.25, 3.75, 3.48, ph + 0.45)
text(s, 9.5, 3.9, 3.0, 0.35, "Ejemplo mínimo", 15, C["muted"], font=MONO)
text(s, 9.5, 4.3, 3.0, 0.4, "n = 2  (N = 4)", 18, font=MONO)
text(s, 9.5, 4.8, 3.0, 0.7, "[[1]] iteración", 26, bold=True)
text(s, 9.5, 5.4, 3.0, 0.7, "[[100 %]] de acierto", 26, bold=True)
chips(s, 0.6, 6.25, ["H", "oráculo", "difusor", "iteración"], size=14)

# ================================================================ 7 curva sin²
s = slide(7, "La curva sin²",
          "Expone B. Idea clave: la probabilidad sube hasta el óptimo y vuelve a bajar; más "
          "iteraciones no siempre es mejor.\nCon n=4 el óptimo es k=3 (96,1 %). Statevector "
          "coincide con la fórmula: error máximo 6,15·10⁻¹⁴.")
pic, pw, ph = picture(s, IMG["curva"], 0.6, 1.9, w=7.35)
text(s, 0.6, 1.9 + ph + 0.12, 7.35, 0.35, "Laboratorio: n=4, estado 1010, curva ideal", 14, C["muted"])
x = 8.4
panel(s, x, 1.9, 4.33, 1.5)
text(s, x + 0.25, 2.05, 3.9, 0.55, "P = sin²((2k+1)θ)", 24, C["cyan"], bold=True, font=MONO)
text(s, x + 0.25, 2.75, 3.9, 0.45, "θ = asin(1/√N)", 18, C["muted"], font=MONO)
keyrow(s, x, 3.7, 4.33, "Sube hasta el [[k óptimo]]: n=4 da k=3 y 96,1 %")
keyrow(s, x, 4.55, 4.33, "Si te pasas, [[baja]]: la amplitud gira de más")
keyrow(s, x, 5.4, 4.33, "[[Statevector]] = fórmula, error máx. {{6,15·10⁻¹⁴}}")

# ================================================================ 8 recorrido
s = slide(8, "Recorrido por la app",
          "Expone C. Idea clave: nueve secciones en la barra, del problema y la ruta guiada "
          "al reto diario y las limitaciones.\n"
          "Señalar la barra de navegación y el panel «Explora la escala» del inicio; "
          "todo corre en local, sin cuenta IBM.")
pic, pw, ph = picture(s, IMG["inicio"], 0.6, 1.85, w=7.9)
k = pw / 2880
panel(s, 0.6 + 240 * k, 1.85 + 140 * k, 2050 * k, 82 * k, fill=None, line="cyan", radius=0.06, lw=2)
text(s, 0.6, 1.85 + ph + 0.1, 7.9, 0.35, "Inicio: navegación y «Explora la escala»", 14, C["muted"])
x = 8.95
text(s, x, 1.78, 1.2, 0.8, "9", 48, C["cyan"], bold=True, font=MONO)
text(s, x + 0.75, 1.98, 3.0, 0.5, "secciones", 22, bold=True)
rows = [("El problema", "para quién"), ("Tu ruta", "pasos guiados"),
        ("Reto diario", "una al día"), ("Laboratorio", "predice y mide"),
        ("Amplitudes", "en vivo"), ("Azar vs Grover", "Monte Carlo"),
        ("Tu contraseña", "entropía"), ("Poscuántica", "Shor y NIST"),
        ("Limitaciones", "lo no probado")]
for i, (a, b) in enumerate(rows):
    y = 2.68 + i * 0.44
    text(s, x, y, 0.4, 0.4, str(i + 1), 16, C["cyan"], font=MONO)
    text(s, x + 0.4, y, 1.8, 0.4, a, 17)
    text(s, x + 2.2, y + 0.02, 1.55, 0.4, b, 14, C["muted"])

# ================================================================ 9 POE
s = slide(9, "Predice → Observa → Explica",
          "Expone C. Idea clave: antes de ejecutar te comprometes con una predicción; luego "
          "comparas con teoría y medida.\nEjemplo de la captura: predijo 70 %, teoría 96,1 %. "
          "Acercarse da XP; equivocarse enseña (metacognición).")
poe = [("1", "Predice", "te comprometes con un número"),
       ("2", "Observa", "teoría y medida, lado a lado"),
       ("3", "Explica", "por qué acertaste o fallaste")]
for i, (num, head, body) in enumerate(poe):
    x = 0.6 + i * 4.11
    panel(s, x, 1.82, 3.91, 0.95)
    text(s, x + 0.25, 1.93, 0.4, 0.4, num, 20, C["cyan"], bold=True, font=MONO)
    text(s, x + 0.65, 1.9, 3.1, 0.4, head, 20, bold=True)
    text(s, x + 0.65, 2.3, 3.15, 0.4, body, 15, C["muted"])
pic, pw, ph = picture(s, IMG["prediccion"], 0.6, 3.0, h=3.3)
arrow(s, 0.6 + pw + 0.2, 3.0 + ph / 2 - 0.17)
picture(s, IMG["poe"], 0.6 + pw + 0.8, 3.0, h=3.3)
text(s, 0.6, 6.42, 12.1, 0.4, "[[POE]]: comprometerse antes activa la [[metacognición]]; acercarse a la teoría da [[XP]].", 16, C["muted"])

# ================================================================ 10 explorador
s = slide(10, "Explorador de amplitudes en vivo",
          "Expone C. Idea clave: mueves k y ves girar el estado y crecer la amplitud del "
          "marcado.\nOráculo invierte el signo, difusor refleja respecto a la media; cada "
          "iteración gira 2θ. Todo en el navegador con fórmulas exactas.")
pic, pw, ph = picture(s, IMG["explorador"], 0.6, 1.85, w=7.7)
text(s, 0.6, 1.85 + ph + 0.1, 7.7, 0.35, "n=4, k=2: barras de amplitud y giro del estado", 14, C["muted"])
x = 8.75
keyrow(s, x, 1.95, 3.98, "[[Oráculo]]: el marcado cambia de signo", h=0.8)
keyrow(s, x, 2.85, 3.98, "[[Difusor]]: todo se refleja respecto a la [[media]]; lo que quedó debajo sube", h=1.2)
keyrow(s, x, 3.95, 3.98, "Cada iteración gira el estado [[2θ]] hacia el marcado", h=0.9)
panel(s, x, 5.2, 3.98, 1.05)
text(s, x + 0.22, 5.3, 3.6, 0.9, "Se calcula en tu navegador con fórmulas exactas, sin servidor.", 16, C["muted"])

# ================================================================ 11 ruido
s = slide(11, "El ruido borra la amplificación",
          "Expone C. Idea clave: con el modelo de ruido FakeFez el éxito cae del 80 % (n=3) "
          "al nivel del azar (n=6).\nCausa: la profundidad pasa de 137 a 4447 y las puertas de "
          "2 qubits de 39 a 1706. Es simulación, no hardware real.")
picture(s, IMG["g_ruido"], 0.45, 1.75, w=7.6, border=False)
x = 8.55
text(s, x, 1.85, 4.2, 0.8, "80 % → azar", 36, C["cyan"], bold=True, font=MONO)
text(s, x, 2.7, 4.18, 1.2, "Más qubits, circuito más [[profundo]]: cada puerta suma error. "
     "Puertas de 2 qubits: 39 con n=3, 1706 con n=6.", 17)
pic, pw, ph = picture(s, IMG["ruido"], x, 4.1, w=4.18)
text(s, x, 4.1 + ph + 0.08, 4.18, 0.35, "App, n=6 con ruido: histograma plano", 14, C["muted"])
chips(s, x, 4.1 + ph + 0.5, ["NISQ", "FakeFez", "profundidad"], size=14)

# ================================================================ 12 IBM real
s = slide(12, "Hacia una computadora cuántica real de IBM",
          "Expone D. Idea clave: el mismo circuito se transpila para el chip y se envía con "
          "SamplerV2.\nTres modos: fake (probado), dry-run e ibm (listos, pendientes de cuenta). "
          "Aún no enviamos ningún job real.")
nodes = [("Tu laptop", "run_ibm_hardware.py"), ("IBM Quantum", "QiskitRuntimeService"),
         ("QPU", "cola y job"), ("Resultados", "counts → JSON")]
nw, ng = 2.75, 0.38
for i, (head, code) in enumerate(nodes):
    x = 0.6 + i * (nw + ng)
    panel(s, x, 1.85, nw, 1.25, line="cyan" if i == 2 else "line")
    text(s, x + 0.18, 1.97, nw - 0.3, 0.45, head, 19, bold=True)
    text(s, x + 0.18, 2.47, nw - 0.3, 0.5, code, 14, C["cyan"], font=MONO)
    if i < 3:
        arrow(s, x + nw + 0.05, 2.3, w=0.28)
for i, lab in enumerate(["transpile, nivel 3", "SamplerV2.run", "job_id"]):
    cx = 0.6 + i * (nw + ng) + nw + ng / 2
    text(s, cx - 1.5, 3.2, 3.0, 0.35, lab, 14, C["muted"], font=MONO, align=PP_ALIGN.CENTER)
modes = [("fake", "FakeFez local, sin red ni cuenta", "probado", "cyan"),
         ("dry-run", "consulta el backend y transpila, no envía", "pendiente de cuenta", "amber"),
         ("ibm", "envía un job a la QPU, consume cuota", "pendiente de cuenta", "amber")]
for i, (m, body, st, col) in enumerate(modes):
    x = 0.6 + i * 4.11
    panel(s, x, 3.75, 3.91, 1.7)
    text(s, x + 0.25, 3.88, 3.4, 0.4, f"--mode {m}", 17, C["text"], bold=True, font=MONO)
    text(s, x + 0.25, 4.35, 3.45, 0.6, body, 15, C["muted"])
    w = chip_w(st, 14)
    panel(s, x + 0.25, 4.95, w, 0.36, fill="panel", line=col, radius=0.18)
    text(s, x + 0.25, 4.95, w, 0.36, st, 14, C[col], font=MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
panel(s, 0.6, 5.68, 12.13, 0.6, fill="panel")
text(s, 0.85, 5.68, 11.7, 0.62, "python -m scripts.run_ibm_hardware --mode ibm --n-qubits 3 --marked 101",
     16, C["cyan"], font=MONO, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.45, 12, 0.4, "Honestidad: el CLI está listo y probado con dobles locales; aún no enviamos ningún job real.", 15, C["amber"])

# ================================================================ 13 contraseñas y PQC
s = slide(13, "Contraseñas y criptografía poscuántica",
          "Expone D. Idea clave: Grover reduce a la mitad los bits efectivos; Shor rompe RSA; "
          "NIST ya publicó reemplazos.\nGenerador local con CSPRNG: 2072 palabras, 11,02 bits "
          "por palabra; meta ≥128 bits ≈ 64 efectivos frente a Grover idealizado.")
cw = 5.86
for i, (lab, col, big, body, img, cap, kw) in enumerate([
    ("Grover: contraseñas", "cyan", "128 → 64 bits",
     "Grover [[reduce a la mitad]] los bits efectivos. Frase de 2072 palabras: 11,02 bits por palabra.",
     "generador", "Generador en el navegador (CSPRNG)", ["entropía", "CSPRNG"]),
    ("Shor: clave pública", "mag", "RSA → ML-KEM",
     "Shor rompe <<RSA y ECC>>. NIST ya publicó reemplazos clásicos: <<ML-KEM>> y <<ML-DSA>>.",
     "pqc", "Tarjeta de la sección Poscuántica", ["ML-KEM", "ML-DSA"]),
]):
    x = 0.6 + i * (cw + 0.41)
    text(s, x, 1.82, 3.4, 0.35, lab, 15, C[col], font=MONO)
    chips(s, x + cw - sum(chip_w(k, 14) + 0.12 for k in kw) + 0.12, 1.78, kw, accent=col, size=14, gap=0.12)
    text(s, x, 2.25, cw, 0.8, big, 36, C[col], bold=True, font=MONO)
    text(s, x, 3.1, cw, 0.95, body, 18)
    pic, pw, ph = picture(s, IMG[img], x, 4.15, w=cw)
    text(s, x, 4.15 + ph + 0.08, cw, 0.35, cap, 14, C["muted"])

# ================================================================ 14 arquitectura
s = slide(14, "Arquitectura y seguridad",
          "Expone D. Idea clave: un solo proceso FastAPI sirve la API y la web; nada de "
          "contraseñas sale del navegador.\nCSP de mismo origen, validación en API y motor; "
          "IBM Quantum solo se usa desde el CLI, aparte del servidor.")
arch = [("Navegador", "TypeScript + SVG", "contraseñas, explorador\ny progreso, en local"),
        ("FastAPI", "un solo proceso", "API + web, validación\ny cabeceras CSP"),
        ("Qiskit / Aer", "FakeFez · NumPy", "circuitos y\nMonte Carlo locales")]
aw, ag = 3.35, 1.04
for i, (head, code, body) in enumerate(arch):
    x = 0.6 + i * (aw + ag)
    panel(s, x, 1.85, aw, 2.0, line="cyan" if i == 1 else "line")
    text(s, x + 0.25, 1.98, aw - 0.5, 0.45, head, 22, bold=True)
    text(s, x + 0.25, 2.47, aw - 0.5, 0.4, code, 15, C["cyan"], font=MONO)
    text(s, x + 0.25, 2.92, aw - 0.5, 0.85, body, 16, C["muted"])
    if i < 2:
        arrow(s, x + aw + 0.17, 2.68, w=0.7, h=0.36, kind=MSO_SHAPE.LEFT_RIGHT_ARROW)
        text(s, x + aw, 3.1, ag, 0.3, "JSON", 14, C["muted"], font=MONO, align=PP_ALIGN.CENTER)
x = 0.6 + 2 * (aw + ag)
s.shapes.add_connector(1, Inches(x + aw / 2), Inches(3.85), Inches(x + aw / 2), Inches(4.2)).line.color.rgb = C["muted"]
panel(s, x, 4.2, aw, 0.75, fill=None, line="muted", dash=True)
text(s, x + 0.25, 4.2, aw - 0.5, 0.75, "IBM Quantum: solo el CLI", 16, C["muted"], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.7, 7.6, 0.4, "Lo que ve el estudiante en la sección de contraseñas:", 16, C["muted"])
picture(s, IMG["privacidad"], 0.6, 5.15, w=12.13)
chips(s, 0.6, 6.0, ["FastAPI", "TypeScript", "CSP default-src 'self'", "privacidad"], size=14)

# ================================================================ 15 validación
s = slide(15, "Validación y limitaciones",
          "Expone D. Idea clave: 72 pruebas Python y 37 de frontend pasan; no corrimos hardware "
          "real ni el piloto con alumnos.\nTodas las cifras son medidas en este repositorio, "
          "con seed fija y reproducibles.")
text(s, 0.6, 1.82, 5.8, 0.35, "Medido y reproducible", 15, C["cyan"], font=MONO)
for i, (num, lab) in enumerate([("72", "pruebas Python\n(pytest)"), ("37", "pruebas de\nfrontend")]):
    x = 0.6 + i * 2.95
    text(s, x, 2.2, 1.2, 1.0, num, 54, C["cyan"], bold=True, font=MONO)
    text(s, x + 1.2, 2.38, 1.6, 0.8, lab, 16, C["muted"])
rows = [("Grover n=4, k=3", "0,953 vs 0,961"), ("Curva vs fórmula", "error 6,15·10⁻¹⁴"),
        ("Monte Carlo 10 bits", "513,3 vs 512,5")]
panel(s, 0.6, 3.45, 5.85, 2.55)
for i, (a, b) in enumerate(rows):
    y = 3.62 + i * 0.78
    text(s, 0.85, y, 2.6, 0.5, a, 17)
    text(s, 3.4, y, 2.9, 0.5, b, 17, C["cyan"], font=MONO)
    if i < 2:
        s.shapes.add_connector(1, Inches(0.85), Inches(y + 0.62), Inches(6.2), Inches(y + 0.62)).line.color.rgb = C["line"]
x = 6.88
text(s, x, 1.82, 5.8, 0.35, "Lo que no demostramos", 15, C["amber"], font=MONO)
lims = ["[[Sin hardware real]] de IBM: el CLI está listo, 0 jobs enviados",
        "[[Sin piloto]] con estudiantes: no afirmamos mejoras de aprendizaje",
        "El oráculo [[conoce la respuesta]]: no rompe contraseñas reales",
        "Ventaja [[cuadrática]], no exponencial"]
for i, t in enumerate(lims):
    y = 2.25 + i * 0.95
    rect(s, x, y + 0.13, 0.11, 0.11, "amber")
    text(s, x + 0.3, y, 5.55, 0.85, t.replace("[[", "").replace("]]", ""), 18)
chips(s, 0.6, 6.3, ["reproducible", "honestidad", "límites"], size=14)

# ================================================================ 16 cierre
s = slide(16, "Impacto EdTech y cierre",
          "Expone A. Idea clave: desde 14 años, 3 minutos al día, con racha y retos: aprender "
          "jugando y midiendo.\nPróximos pasos: piloto con estudiantes, primer job real en IBM, "
          "uso en aula. Cerrar con la invitación a probarlo.")
stats = [("14+", "años: no exige\nálgebra lineal"), ("3 min", "al día: reto,\npredicción y circuito"),
         ("racha", "XP y niveles\npara volver")]
for i, (big, lab) in enumerate(stats):
    y = 1.85 + i * 1.32
    text(s, 0.6, y, 2.4, 0.9, big, 40, C["cyan"], bold=True, font=MONO)
    text(s, 2.75, y + 0.12, 2.5, 0.85, lab, 16, C["muted"])
pic, pw, ph = picture(s, IMG["reto"], 5.6, 1.9, w=7.13)
text(s, 5.6, 1.9 + ph + 0.08, 7.13, 0.35, "Reto diario: una pregunta nueva cada día", 14, C["muted"])
y = 4.85
text(s, 5.6, y, 3, 0.4, "Próximos pasos", 18, bold=True)
for xx, num, t in [(5.6, "1", "Piloto con alumnos"), (8.15, "2", "Job real en IBM"), (10.55, "3", "Uso en clase")]:
    text(s, xx, y + 0.45, 0.35, 0.4, num, 17, C["cyan"], font=MONO)
    text(s, xx + 0.32, y + 0.45, 2.1, 0.4, t, 17)
panel(s, 0.6, 5.95, 12.13, 0.78, fill="cyan", line=None)
text(s, 0.9, 5.95, 11.6, 0.78, "Pruébalo: predice antes de ejecutar.", 24, C["ink"], bold=True,
     anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print(f"OK {OUT} ({len(prs.slides)} slides)")
